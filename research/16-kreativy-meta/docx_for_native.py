"""Docx для носителя языка: на каждый креатив — картинка, текст на узбекском, исходный русский.

Узбекский — финальные JSON (то, что стоит на макетах), русский — исходники creatives_ru*.json
и 15-strategiya-targeta-meta/04-texts-static.json (карусели). Картинки — из папки на рабочем столе.

python3 docx_for_native.py
"""
import json
import re
import shutil
import subprocess
from math import ceil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
DESK = Path.home() / "Desktop" / "Креативы УЗ — Meta"
STATIC_RU = ROOT.parent / "15-strategiya-targeta-meta" / "04-texts-static.json"
CACHE = ROOT / "render" / "docx_native"
OUT = ROOT / "Тексты креативов UZ-RU — для Силбека.docx"
DESK_OUT = Path.home() / "Desktop" / OUT.name
FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
GREY = RGBColor(0x6B, 0x6B, 0x6B)
COL_W = (Cm(3.3), Cm(6.85), Cm(6.85))

LABEL = {
    "headline": "Заголовок", "subline": "Подзаголовок", "course_name": "Название курса",
    "cta": "Призыв", "sticker_title": "Шапка стикера", "question": "Вопрос", "answer": "Ответ",
    "bullets": "Плашка", "price_old": "Старая цена (зачёркнута)", "price_new": "Новая цена",
    "caption": "Крупная надпись", "line1": "Строка 1", "line2": "Строка 2", "line3": "Строка 3",
    "title_line1": "Плашка 1", "title_line2": "Плашка 2", "today": "Справа сверху", "free": "Крупно",
    "top": "Над мемом", "in_meme": "Внутри мема", "bottom": "Под мемом",
}
SKIP = {"payment", "kw4", "kw5", "lines", "logo_word_white", "logo_word_box"}
NIKITA = ("Подпись у фото", "Nikita / kurs muallifi", "Никита / автор курса")

GROUPS = [
    ("01 Инста — эксперт на фото", ["I1-1", "I1-2", "I1-3", "I1-3B", "I1-4"]),
    ("02 Инста — вопрос-ответ", ["I2-1", "I2-2", "I2-3", "I2-4"]),
    ("03 Инста — текст и фон", ["I3-1", "I3-2", "I3-3"]),
    ("04 Мемы", ["M1", "M2", "M3", "M4", "M5", "M6", "M7"]),
    ("05 Видео", ["V1", "V2", "V3-V4"]),
    ("06 Красивые — карусели 1x1", ["K2-1", "K2-1b", "K2-3", "K2-4"]),
    ("08 Инста — короткие тексты (A-B)", ["I1-1S", "I1-2S", "I2-1S", "I2-2S", "I2-3S", "I2-4S", "I3-1S", "I3-3S"]),
]

# Кадры для видео: секунды, колонки сетки
FRAMES = {
    "I3-2": ([0.8, 3.4, 5.0, 9.0], 4),
    "M2": ([20.5], 1),
    "V1": ([1.0, 3.0, 5.0, 7.0, 9.0, 11.0, 13.0, 15.5], 4),
    "V2": ([1.2, 2.5, 4.0, 5.5, 6.8, 8.2, 9.5, 11.5], 4),
}

V34_ROWS = [
    ("Начало", "Bu rolikni AI yaratdi", "Этот ролик сделал ИИ"),
    ("Про цену", "Bunaqasini mijozlarga $100 dan sotamiz", "Такие мы продаём клиентам по $100"),
    ("Плашка (только V3)", "TOP-3", "ТОП-3"),
    ("Без знаний", "Dizayn, montaj va dasturlashni bilmasdan qilamiz", "И делаем без знаний дизайна, монтажа и программирования"),
    ("Плашки", "dizayn · montaj · dasturlash", "дизайн · монтаж · программирование"),
    ("Плашка", "AI-MONTAJ", "AI-МОНТАЖ"),
    ("Время", "Bitta rolikka — ikki soatcha", "На один ролик — пару часов"),
    ("Расчёт", "Hisob-kitob, va’da emas · 10 × $100", "Расчёт, не обещание · 10 × $100"),
    ("Вопрос", "Siz ham onlayn AI-mutaxassis kasbini egallab, $1000 ga chiqmoqchimisiz?",
     "Тоже хочешь освоить онлайн-профессию AI-специалиста и выйти на $1000?"),
    ("Плашка", "Bepul mini-kurs", "Бесплатный мини-курс"),
    ("Галочки", "mutlaqo noldan · dasturlashni bilmasdan · sarmoyasiz",
     "с полного нуля · без знаний программирования · без вложений"),
    ("Стена надписей", "BEPUL · kun oxirigacha", "БЕСПЛАТНО · до конца дня"),
    ("Дедлайн (только V3)", "Kun oxirigacha olishga ulguring, hozircha bepul.", "Успей забрать до конца дня, пока он бесплатный."),
    ("Цена", "0 so‘m · Keyin — 3 mln so‘m", "0 сум · Потом — 3 млн сум"),
    ("Призыв", "Pastdagi tugmani bosing va bepul kursni ko‘ring", "Переходи по кнопке ниже и смотри бесплатный курс"),
    ("Карточка курса", "Bepul mini-kurs · Onlayn AI-mutaxassis kasbi · bepul",
     "Бесплатный мини-курс · Онлайн-профессия AI-специалиста · бесплатно"),
    ("Кнопка (только V3)", "Bepul ko‘rish", "Смотреть бесплатно"),
    ("Финал", "Bosing", "Кликай"),
]

INTRO = [
    "Это тексты на рекламных картинках и видео онлайн-школы (Instagram и Facebook). "
    "Узбекский перевод сделан с помощью ИИ, живой носитель его ещё не читал.",
    "На каждый креатив: картинка → текст на узбекском (как сейчас на макете) → русский (что хотим сказать).",
]
ASKS = [
    "Проверьте узбекский: звучит ли естественно, нет ли ошибок, совпадает ли смысл с русским.",
    "Правьте прямо в колонке «Узбекский». Если всё хорошо — ничего не меняйте.",
    "Латиница; o‘ и g‘ — с апострофом ‘.",
    "Обращение — на «siz». Русский оригинал написан на «ты», но в узбекской рекламе принято «siz».",
    "Длина: текст должен поместиться на картинку — новая версия не длиннее нынешней больше чем на 20%.",
    "Цифры, суммы и цены не меняем.",
]


# ---------- данные ----------

def load():
    ru = {c["id"]: c["fields_ru"] for c in json.loads((ROOT / "creatives_ru.json").read_text())["creatives"]}
    ru |= {c["id"]: c["fields_ru"] for c in json.loads((ROOT / "creatives_ru_short.json").read_text())["creatives"]}
    uz = {c["id"]: c["fields_uz"] for c in json.loads((ROOT / "creatives_uz.json").read_text())["creatives"]}
    uz |= {c["id"]: c["fields_uz"] for c in json.loads((ROOT / "creatives_uz_short.json").read_text())["creatives"]}
    car_uz = {c["id"]: c for c in json.loads((ROOT / "creatives_uz_type2.json").read_text())["texts"]}
    car_ru = find_ids(json.loads(STATIC_RU.read_text()), set(car_uz))
    missing = set(car_uz) - set(car_ru)
    if missing:
        raise KeyError(f"нет русского текста каруселей: {missing}")
    return ru, uz, car_uz, car_ru


def find_ids(obj, ids):
    found = {}

    def walk(o):
        if isinstance(o, dict):
            if o.get("id") in ids and o["id"] not in found and "cards" in o:
                found[o["id"]] = o
            for v in o.values():
                walk(v)
        elif isinstance(o, list):
            for v in o:
                walk(v)
    walk(obj)
    return found


def field_rows(fu: dict, fr: dict) -> list:
    rows = []
    for key, val in fu.items():
        if key in SKIP:
            continue
        label = LABEL.get(key, key)
        ru_val = fr.get(key)
        if isinstance(val, list):
            ru_list = ru_val if isinstance(ru_val, list) else []
            for i, v in enumerate(val):
                n = f" {i + 1}" if len(val) > 1 else ""
                rows.append((label + n, v, ru_list[i] if i < len(ru_list) else "—"))
        else:
            rows.append((label, val, ru_val or "—"))
    return rows


def plain(text: str) -> str:
    """Убирает разметку вёрстки M2: ~~зачёркнуто~~, ==маркер==, [[обведено]]."""
    return re.sub(r"~~|==|\[\[|\]\]", "", text)


def static_rows(cid: str, fu: dict, fr: dict) -> list:
    if cid == "M1":
        head = [("Логотип (слово + плашка)", f"{fu['logo_word_white']} + {fu['logo_word_box']} (читается «Yarating»)",
                 f"{fr['logo_word_white']} + {fr['logo_word_box']} (читается «Научись»)")]
        return head + field_rows(fu, fr)
    if cid == "M2":
        return [(f"Строка {i + 1}", plain(b["text"]), plain(r["text"])) for i, (b, r) in enumerate(zip(fu["blocks"], fr["blocks"]))]
    if cid == "V1":
        s, sr = fu["scenes"], fr["scenes"]
        return [
            ("Сцена 1", s[0], sr[0]),
            ("Сцена 2", s[1], sr[1]),
            ("Сцена 2, подписи телефонов", "SAYT · BOT · REELS", "САЙТ · БОТ · REELS"),
            ("Сцена 3", s[2], sr[2]),
            ("Сцена 3, плашки (зачёркиваются)", " · ".join(fu["chips"]), " · ".join(fr["chips"])),
            ("Сцена 4", s[3], sr[3]),
            ("Сцена 5", s[4], sr[4]),
            ("Сцена 6, галочки", " · ".join(fu["bullets"]), " · ".join(fr["bullets"])),
            ("Сцена 7", fu["cta_scenes"][0], fr["cta_scenes"][0]),
            ("Сцена 7, кнопка", fu["button"], fr["button"]),
            ("Сцена 8, под таймером", fu["cta_scenes"][1], fr["cta_scenes"][1]),
            ("Сцена 9, финал", fu["cta_scenes"][2], fr["cta_scenes"][2]),
        ]
    if cid == "V2":
        return [
            ("Дата над часами", "Yakshanba, 11-oktabr", "Воскресенье, 11 октября"),
            ("Время в уведомлениях", "hozir", "сейчас"),
            ("Уведомление 1 (Uzum Bank)", fu["notif_bank"], fr["notif_bank"]),
            ("Уведомление 2, от кого", "Dilshod (mijoz)", "Дилшод (клиент)"),
            ("Уведомление 2", fu["notif_client"], fr["notif_client"]),
            ("Уведомление 3 (Sardor)", fu["notif_tg"], fr["notif_tg"]),
            ("Крупная надпись", fu["caption"], fr["caption"]),
            ("Кнопка", fu["cta"], fr["cta"]),
        ]
    rows = field_rows(fu, fr)
    if cid.startswith("I1-"):
        rows.append(NIKITA)
    if cid == "I1-3B":
        rows.append(("Скрин оплаты", "не переводим — настоящий скрин банка", "—"))
    return rows


def cards_of(c: dict, sfx: str) -> list:
    cards = [dict(x) for x in c["cards"]]
    head, sub = c[f"headline{sfx}"], c[f"subline{sfx}"]
    if cards[0][f"title{sfx}"] != head:
        cards.insert(0, {f"title{sfx}": head, f"text{sfx}": sub})
    return cards


def carousel_rows(cu: dict, cr: dict) -> list:
    uz_cards, ru_cards = cards_of(cu, "_uz"), cards_of(cr, "")
    rows = [("Карточка 1, плашка сверху", cu["badge_uz"], cr["badge"])]
    for n, (u, r) in enumerate(zip(uz_cards, ru_cards), start=1):
        rows.append((f"Карточка {n}, заголовок", u["title_uz"], r["title"]))
        rows.append((f"Карточка {n}, текст", u["text_uz"], r["text"]))
        if u.get("small_uz"):
            rows.append((f"Карточка {n}, мелкая строка", u["small_uz"], r.get("small", "—")))
    rows.append((f"Карточка {len(uz_cards)}, кнопка", cu["cta_uz"], cr["cta"]))
    return rows


# ---------- картинки ----------

def desk_file(cid: str, ext: str) -> Path:
    hits = sorted(DESK.glob(f"*/{cid} *4x5.{ext}"))
    if not hits:
        raise FileNotFoundError(f"{cid}: нет файла 4x5.{ext} в {DESK}")
    return hits[0]


def grid(images: list, cols: int, name: str, labels=None, cell_w=320, gap=10) -> Path:
    h = round(images[0].height * cell_w / images[0].width)
    lab_h = 30 if labels else 0
    rows = ceil(len(images) / cols)
    canvas = Image.new("RGB", (cols * cell_w + (cols + 1) * gap, rows * (h + lab_h) + (rows + 1) * gap), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(FONT, 18)
    for i, im in enumerate(images):
        x = gap + (i % cols) * (cell_w + gap)
        y = gap + (i // cols) * (h + lab_h + gap)
        canvas.paste(im.resize((cell_w, h)), (x, y))
        if labels:
            draw.text((x + 4, y + h + 5), labels[i], fill=(90, 90, 90), font=font)
    out = CACHE / f"{name}.jpg"
    canvas.save(out, quality=86)
    return out


def video_strip(mp4: Path, times: list, cols: int, name: str) -> Path:
    tiles = []
    for i, t in enumerate(times):
        p = CACHE / f"{name}_{i}.png"
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-ss", str(t), "-i", str(mp4),
                        "-frames:v", "1", str(p)], check=True)
        tiles.append(Image.open(p).convert("RGB"))
    if len(tiles) == 1:
        return still(tiles[0], name)
    return grid(tiles, cols, name, labels=[f"{t:g} с" for t in times])


def still(im: Image.Image, name: str, max_w=900) -> Path:
    im = im.convert("RGB")
    im.thumbnail((max_w, max_w * 2))
    out = CACHE / f"{name}.jpg"
    im.save(out, quality=86)
    return out


def picture(cid: str) -> tuple:
    """Картинка, ширина в см и подпись о файле."""
    if cid.startswith("K"):
        folder = DESK / "06 Красивые — карусели 1x1" / f"{cid} — карусель"
        cards = sorted(folder.glob(f"{cid}_card*_1x1.png"))
        img = grid([Image.open(c).convert("RGB") for c in cards], 3, cid,
                   labels=[f"Карточка {i + 1}" for i in range(len(cards))])
        return img, 14.5, f"06 Красивые — карусели 1x1 / {folder.name} ({len(cards)} карточек)"
    if cid == "V3-V4":
        v3 = desk_file("V3", "mp4")
        v4 = desk_file("V4", "mp4")
        img = video_strip(v4, [1.0, 4.5, 8.5, 12.5, 16.5, 21.0, 26.0, 31.5], 4, "V4")
        return img, 16, f"05 Видео / {v3.name}, {v4.name}, V4б (кадры — из V4)"
    if cid in FRAMES:
        mp4 = desk_file(cid, "mp4")
        times, cols = FRAMES[cid]
        width = 7 if len(times) == 1 else 16
        return video_strip(mp4, times, cols, cid), width, f"{mp4.parent.name} / {mp4.stem.rsplit(' — ', 1)[0]} (видео, 4x5 и 9x16)"
    png = desk_file(cid, "png")
    return still(Image.open(png), cid), 7, f"{png.parent.name} / {png.stem.rsplit(' — ', 1)[0]} (4x5 и 9x16 — текст одинаковый)"


# ---------- документ ----------

def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def fill_cell(cell, text: str, *, size=10, bold=False, color=None):
    parts = [p for p in text.split(" · ")] if " · " in text else [text]
    cell.paragraphs[0].text = ""
    for i, part in enumerate(parts):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(part)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color


def text_table(doc, rows: list):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    head = table.rows[0]
    repeat_header(head)
    for cell, title in zip(head.cells, ("Где", "Узбекский (сейчас на макете)", "Русский (что хотим сказать)")):
        fill_cell(cell, title, bold=True)
        shade(cell, "EDEDED")
    for where, uz, ru in rows:
        cells = table.add_row().cells
        fill_cell(cells[0], where, size=9, color=GREY)
        fill_cell(cells[1], uz)
        fill_cell(cells[2], ru)
    for row in table.rows:
        for cell, w in zip(row.cells, COL_W):
            cell.width = w


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        if name != "Normal":
            st.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13)


def creative(doc, cid: str, rows: list, title: str, new_page: bool):
    img, width, note = picture(cid)
    h = doc.add_heading(f"{cid} · {title}", level=2)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.page_break_before = new_page
    p = doc.add_paragraph()
    r = p.add_run(f"Файл: {note}")
    r.font.size, r.font.color.rgb = Pt(9), GREY
    p.paragraph_format.keep_with_next = True
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(img), width=Cm(width))
    text_table(doc, rows)


def title_of(cid: str, ru: dict) -> str:
    if cid == "V3-V4":
        return "Монтаж как в референсе (V3) и на HyperFrames (V4, V4б) — тексты общие"
    if cid.startswith("K"):
        return "карусель"
    f = ru.get(cid, {})
    for key in ("headline", "question", "top", "title_line1"):
        if f.get(key) and cid not in ("M1", "M2", "V1", "V2"):
            return f[key]
    file = desk_file(cid, "mp4" if cid in FRAMES else "png")
    return file.stem.rsplit(" — ", 1)[0].split(" ", 1)[1]


def main():
    CACHE.mkdir(parents=True, exist_ok=True)
    ru, uz, car_uz, car_ru = load()
    doc = Document()
    setup(doc)
    doc.add_heading("Тексты креативов: узбекский и русский", level=0)
    for line in INTRO:
        doc.add_paragraph(line)
    doc.add_paragraph("Что нужно:").runs[0].bold = True
    for line in ASKS:
        doc.add_paragraph(line, style="List Bullet")

    count = 0
    for folder, ids in GROUPS:
        g = doc.add_heading(folder, level=1)
        g.paragraph_format.page_break_before = True
        g.paragraph_format.keep_with_next = True
        for i, cid in enumerate(ids):
            if cid == "V3-V4":
                rows = V34_ROWS
            elif cid.startswith("K"):
                key = cid[1:]
                rows = carousel_rows(car_uz[key], car_ru[key])
            else:
                rows = static_rows(cid, uz[cid], ru[cid])
            creative(doc, cid, rows, title_of(cid, ru), new_page=i > 0)
            count += 1

    doc.save(OUT)
    shutil.copy2(OUT, DESK_OUT)
    print(f"{count} креативов → {OUT}\nкопия: {DESK_OUT}")


if __name__ == "__main__":
    main()
