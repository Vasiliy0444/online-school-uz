"""Правки носителя языка: docx → исходники креативов, проверка, выгрузка на перевод.

Формат docx — как вернул переводчик 15.10.2026: на креатив заголовок «I1-2 · …» (Heading 2)
и таблица «Русский (что хотим сказать) | Узбекский (сейчас на креативе) | Узбекский (как правильно на креативе)».
Изменённые ячейки третьей колонки выделены жёлтым.

python3 native_fixes.py diff  <docx>                        # что поменял носитель
python3 native_fixes.py apply <docx> [--all]                # вписать правки в creatives_uz*.json
                                                            # (--all — и правки без жёлтой заливки)
python3 native_fixes.py check <docx>                        # третья колонка = исходники, иначе код 1
python3 native_fixes.py export I1-3,I1-3B,I3-2 <out.docx>   # docx на перевод в том же формате

V3/V4 берут текст не из JSON: apply их только перечисляет (править text.ts и index.html руками),
check сверяет их изменённые строки с текстом этих файлов.
"""
import html
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

import docx_for_native as N

ROOT = Path(__file__).resolve().parent
SOURCES = {
    "main": ROOT / "creatives_uz.json",
    "short": ROOT / "creatives_uz_short.json",
    "carousel": ROOT / "creatives_uz_type2.json",
}
V34_FILES = (ROOT / "montage-remotion" / "src" / "montage" / "text.ts", ROOT / "v4-hyperframes" / "index.html")
RENDER = ROOT / "render" / "uz"
RENDER_SHORT = ROOT / "render" / "uz_short"
SKIP_KEYS = {"id", "type", "n", "notes_uz", "len_check", "payment"}
NEUTRAL_FILL = {"", "AUTO", "FFFFFF", "EDEDED"}
HEAD = ("Русский (что хотим сказать)", "Узбекский (сейчас на креативе)", "Узбекский (как правильно на креативе)")
COL_W = (N.Cm(5.1), N.Cm(5.95), N.Cm(5.95))

INTRO = [
    "Это тексты на рекламных картинках и видео онлайн-школы (Instagram и Facebook). "
    "Эти креативы не вошли в проверку от 15.10.2026.",
    "На каждый креатив: картинка → таблица из трёх колонок: русский (что хотим сказать) → "
    "узбекский (как сейчас на креативе) → узбекский (как правильно написать на креативе).",
]
ASKS = [
    "Третья колонка — готовый текст для креатива. Сейчас в ней стоит текущий текст.",
    "Если строка верная — не трогайте. Если нет — исправьте третью колонку и выделите ячейку жёлтым.",
    "Латиница; o‘ и g‘ — с апострофом ‘.",
    "Обращение — на «siz». Русский оригинал написан на «ты», но в узбекской рекламе принято «siz».",
    "Длина: исправленный текст не длиннее нынешнего больше чем на 20%, иначе не поместится на картинку.",
    "Цифры, суммы и цены не меняем.",
]


@dataclass(frozen=True)
class Row:
    label: str
    ru: str
    now: str
    fix: str
    marked: bool


# ---------- чтение docx ----------

def cell_text(cell) -> str:
    """Абзацы ячейки через « · » — так их разбил docx_for_native; пустые (лишний Enter) не считаем."""
    return " · ".join(p.text for p in cell.paragraphs if p.text.strip())


def is_marked(cell) -> bool:
    """Ячейка выделена: заливка ячейки или маркер на тексте."""
    pr = cell._tc.tcPr
    shd = pr.find(qn("w:shd")) if pr is not None else None
    fill = (shd.get(qn("w:fill")) or "").upper() if shd is not None else ""
    highlight = any(r._r.rPr is not None and r._r.rPr.find(qn("w:highlight")) is not None
                    for p in cell.paragraphs for r in p.runs)
    return fill not in NEUTRAL_FILL or highlight


def table_rows(table: Table, cid: str, name: str) -> list:
    head = tuple(c.text.strip() for c in table.rows[0].cells)
    if len(head) != 3 or "как правильно" not in head[2]:
        raise SystemExit(f"{name}: у {cid} другая таблица {head}. Нужен формат «{' | '.join(HEAD)}»")
    rows = []
    for r in table.rows[1:]:
        c = r.cells
        label, *ru = [p.text for p in c[0].paragraphs]
        rows.append(Row(label, " · ".join(ru), cell_text(c[1]), cell_text(c[2]), is_marked(c[2])))
    return rows


def parse(path: Path) -> dict:
    """{id креатива: [Row, …]} в порядке документа."""
    doc = Document(str(path))
    if any(True for _ in doc.element.body.iter(qn("w:ins"), qn("w:del"))):
        raise SystemExit(f"{path.name}: в файле режим правок Word. Примите все изменения "
                         "(Рецензирование → Принять все), сохраните и запустите снова.")
    out, cur = {}, None
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, doc)
            if p.style.name == "Heading 2":
                cur = p.text.split(" · ")[0].strip()
                out[cur] = []
        elif tag == "tbl" and cur:
            out[cur] = out[cur] + table_rows(Table(el, doc), cur, path.name)
    return out


def changes(doc: dict) -> list:
    return [(cid, r) for cid, rows in doc.items() for r in rows if r.now != r.fix]


# ---------- исходники ----------

def load_sources() -> dict:
    return {k: json.loads(p.read_text(encoding="utf-8")) for k, p in SOURCES.items()}


def locate(data: dict, cid: str):
    """(файл, индекс) креатива в исходниках или None — для V3/V4, у которых текст не в JSON."""
    if cid.startswith("K"):
        hits = [i for i, t in enumerate(data["carousel"]["texts"]) if t["id"] == cid[1:]]
        return ("carousel", hits[0]) if hits else None
    for key in ("main", "short"):
        hits = [i for i, c in enumerate(data[key]["creatives"]) if c["id"] == cid]
        if hits:
            return key, hits[0]
    return None


def replaced(obj, old: str, new: str):
    """Копия obj, где каждая строка, равная old, заменена на new, и число замен."""
    if isinstance(obj, str):
        return (new, 1) if obj == old else (obj, 0)
    if isinstance(obj, list):
        parts = [replaced(v, old, new) for v in obj]
        return [v for v, _ in parts], sum(n for _, n in parts)
    if isinstance(obj, dict):
        parts = {k: (v, 0) if k in SKIP_KEYS else replaced(v, old, new) for k, v in obj.items()}
        return {k: v for k, (v, _) in parts.items()}, sum(n for _, n in parts.values())
    return obj, 0


def with_fix(data: dict, where: tuple, old: str, new: str):
    key, i = where
    if key == "carousel":
        items = data[key]["texts"]
        item, n = replaced(items[i], old, new)
        return {**data, key: {**data[key], "texts": items[:i] + [item] + items[i + 1:]}}, n
    items = data[key]["creatives"]
    fields, n = replaced(items[i]["fields_uz"], old, new)
    item = {**items[i], "fields_uz": fields}
    return {**data, key: {**data[key], "creatives": items[:i] + [item] + items[i + 1:]}}, n


def save(key: str, part: dict) -> None:
    path = SOURCES[key]
    tail = "\n" if path.read_text(encoding="utf-8").endswith("\n") else ""
    path.write_text(json.dumps(part, ensure_ascii=False, indent=2) + tail, encoding="utf-8")


def flat(path: Path) -> str:
    """Текст файла V3/V4 одной строкой: строковые литералы text.ts или index.html без тегов."""
    raw = path.read_text(encoding="utf-8")
    if path.suffix == ".ts":
        raw = " ".join(re.findall(r'"((?:[^"\\]|\\.)*)"', raw))
    else:
        raw = html.unescape(re.sub(r"<[^>]+>", " ", raw))
    return re.sub(r"\s+", " ", raw)


# ---------- команды ----------

def diff(path: Path) -> int:
    doc = parse(path)
    found = changes(doc)
    for cid, r in found:
        flag = "" if r.marked else "  (не выделено жёлтым)"
        print(f"{cid} · {r.label}{flag}\n  было:  {r.now}\n  стало: {r.fix}")
    idle = [(cid, r.label) for cid, rows in doc.items() for r in rows if r.marked and r.now == r.fix]
    for cid, label in idle:
        print(f"{cid} · {label}: выделено, но текст тот же")
    print(f"креативов в файле: {len(doc)}, правок: {len(found)}")
    return 0


def refuse_unmarked(found: list) -> None:
    """Правка без жёлтой заливки похожа на случайную (автозамена, лишний пробел) — без --all не вносим."""
    unmarked = [f"  {cid} · {r.label}: «{r.now}» → «{r.fix}»" for cid, r in found if not r.marked]
    if unmarked:
        raise SystemExit("Изменено, но не выделено жёлтым:\n" + "\n".join(unmarked)
                         + "\nЕсли правки настоящие — запустите с --all. Ничего не записано.")


def apply(path: Path, include_unmarked: bool = False) -> int:
    found = changes(parse(path))
    if not include_unmarked:
        refuse_unmarked(found)
    data, manual, log = load_sources(), [], []
    for cid, r in found:
        where = locate(data, cid)
        if where is None:
            manual.append((cid, r))
            continue
        data, n = with_fix(data, where, r.now, r.fix)
        if n == 0 and with_fix(data, where, r.fix, r.fix)[1]:
            log.append(f"{cid} · {r.label}: уже внесено")
            continue
        if n == 0:
            raise SystemExit(f"{cid} · {r.label}: в исходниках нет «{r.now}». Ничего не записано.")
        log.append(f"{cid} · {r.label}: {n} зам.")
    before = load_sources()
    for key in SOURCES:
        if data[key] != before[key]:
            save(key, data[key])
            print("записан", SOURCES[key].name)
    print("\n".join(log))
    for cid, r in manual:
        print(f"РУКАМИ {cid} · {r.label}: «{r.now}» → «{r.fix}» ({', '.join(p.name for p in V34_FILES)})")
    return 0


def expected_rows(cid: str, src: tuple):
    """Строки креатива, как их собирает docx_for_native из текущих исходников; None — для V3/V4."""
    ru, uz, car_uz, car_ru = src
    if cid.startswith("K"):
        return N.carousel_rows(car_uz[cid[1:]], car_ru[cid[1:]])
    if cid in uz:
        return N.static_rows(cid, uz[cid], ru[cid])
    return None


def check_v34(cid: str, rows: list) -> list:
    texts = [flat(p) for p in V34_FILES]
    bad = []
    for r in rows:
        if r.now != r.fix and (not any(r.fix in t for t in texts) or any(r.now in t for t in texts)):
            bad.append(f"{cid} · {r.label}: в text.ts / index.html нет «{r.fix}» или осталось «{r.now}»")
    return bad


def check(path: Path) -> int:
    src, bad = N.load(), []
    for cid, rows in parse(path).items():
        exp = expected_rows(cid, src)
        if exp is None:
            bad += check_v34(cid, rows)
            continue
        if len(exp) != len(rows):
            bad.append(f"{cid}: строк в файле {len(rows)}, в исходниках {len(exp)}")
            continue
        bad += [f"{cid} · {r.label}: в файле «{r.fix}», в исходниках «{e[1]}»"
                for r, e in zip(rows, exp) if r.fix != e[1]]
    print("\n".join(bad) if bad else "всё совпадает: третья колонка = исходники")
    return 1 if bad else 0


# ---------- выгрузка на перевод ----------

def ru_cell(cell, label: str, ru: str) -> None:
    cell.paragraphs[0].text = ""
    run = cell.paragraphs[0].add_run(label)
    run.font.size, run.font.color.rgb = Pt(9), N.GREY
    for part in ru.split(" · "):
        para = cell.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        para.add_run(part).font.size = Pt(10)


def text_table(doc, rows: list) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    N.repeat_header(table.rows[0])
    for cell, title in zip(table.rows[0].cells, HEAD):
        N.fill_cell(cell, title, bold=True)
        N.shade(cell, "EDEDED")
    for label, uz, ru in rows:
        cells = table.add_row().cells
        ru_cell(cells[0], label, ru)
        N.fill_cell(cells[1], uz)
        N.fill_cell(cells[2], uz)
    for row in table.rows:
        for cell, w in zip(row.cells, COL_W):
            cell.width = w


def picture(cid: str) -> tuple:
    if cid.startswith("K"):
        cards = sorted(RENDER.glob(f"{cid}_card*.png"), key=lambda p: int(p.stem.split("card")[1]))
        img = N.grid([Image.open(c).convert("RGB") for c in cards], 3, cid,
                     labels=[f"Карточка {i + 1}" for i in range(len(cards))])
        return img, 14.5, f"карусель, {len(cards)} карточек 1x1"
    if cid.endswith("S"):
        return N.still(Image.open(RENDER_SHORT / f"{cid}_feed.png"), cid), 7, "4x5 и 9x16 — текст одинаковый"
    if cid in N.FRAMES:
        times, cols = N.FRAMES[cid]
        img = N.video_strip(RENDER / f"{cid}_feed.mp4", times, cols, cid)
        return img, 7 if len(times) == 1 else 16, "видео, 4x5 и 9x16"
    return N.still(Image.open(RENDER / f"{cid}_feed.png"), cid), 7, "4x5 и 9x16 — текст одинаковый"


def title(cid: str, ru: dict) -> str:
    f = ru.get(cid, {})
    return next((f[k] for k in ("headline", "question", "top", "title_line1") if f.get(k)), cid)


def export(ids: list, out: Path) -> int:
    N.CACHE.mkdir(parents=True, exist_ok=True)
    src = N.load()
    doc = Document()
    N.setup(doc)
    doc.add_heading(f"Тексты креативов на перевод: {', '.join(ids)}", level=0)
    for line in INTRO:
        doc.add_paragraph(line)
    doc.add_paragraph("Как заполнять:").runs[0].bold = True
    for line in ASKS:
        doc.add_paragraph(line, style="List Bullet")
    for i, cid in enumerate(ids):
        rows = expected_rows(cid, src)
        if rows is None:
            raise SystemExit(f"{cid}: нет в creatives_uz*.json")
        img, width, note = picture(cid)
        h = doc.add_heading(f"{cid} · {title(cid, src[0])}", level=2)
        h.paragraph_format.page_break_before = i > 0
        h.paragraph_format.keep_with_next = True
        p = doc.add_paragraph()
        r = p.add_run(f"Файл: {note}")
        r.font.size, r.font.color.rgb = Pt(9), N.GREY
        pic = doc.add_paragraph()
        pic.alignment = N.WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(str(img), width=N.Cm(width))
        text_table(doc, rows)
    doc.save(str(out))
    print(f"{len(ids)} креативов → {out}")
    return 0


def main(argv: list) -> int:
    if len(argv) < 3 or argv[1] not in ("diff", "apply", "check", "export"):
        print(__doc__)
        return 2
    if argv[1] == "export":
        return export(argv[2].split(","), Path(argv[3]))
    if argv[1] == "apply":
        return apply(Path(argv[2]), include_unmarked="--all" in argv[3:])
    return {"diff": diff, "check": check}[argv[1]](Path(argv[2]))


if __name__ == "__main__":
    sys.exit(main(sys.argv))
