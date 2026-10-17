#!/usr/bin/env python3
"""Сверка одной вкладки с Python-расчётом. python3 verify_single.py <xlsx> <addr.json>"""
import json, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).parent))
import single_model as M

S1K = ["reg", "part", "peak", "orders", "orders_sum", "autopay", "autopay_sum", "prepay", "op_leads", "op_leads_sum",
       "op_bk_plan", "op_or_plan", "web_plan", "op_paid", "op_sum", "roi", "plan", "tg_sub_n"]
S2K = ["from_s1", "buyers", "base", "reg", "part", "autopay", "bk", "orders", "op_leads", "op_bk_plan", "op_or_plan",
       "dozhim_plan", "op_total", "op_sum", "total", "plan"]
TK = ["rev1", "rev2", "op_plan", "launch", "students", "own", "now", "later", "losses", "net", "traffic", "costs",
      "profit", "payback", "roas", "profit_rec", "start_cash", "ok"]

def verify(path, addr_path, s1=M.S1, s2=M.S2, tot=M.TOT, quiet=False):
    cv1, cv2 = M.blended_rates()          # ставки в книге — статичные, посчитаны при сборке по стартовым цифрам
    tw = M.compute(s1, s2, tot, cv1, cv2)
    addr = json.loads(Path(addr_path).read_text())
    ws = load_workbook(path, data_only=True)[M.__dict__.get("SHEET", "Модель продаж - гибридная")]
    problems = 0
    def chk(key, exp):
        nonlocal problems
        col, row = addr[key]
        got = ws[f"{col}{row}"].value
        if isinstance(exp, str):
            ok = got == exp
        else:
            ok = got is not None and not isinstance(got, str) and abs(float(got) - exp) <= max(0.01, abs(exp) * 1e-9)
        if not ok:
            problems += 1
            print(f"MISMATCH {key}: book={got} twin={exp}")
    for n in ("min", "mid", "max"):
        for k in S1K:
            chk(f"S1_{n}:{k}", tw[f"s1_{n}"][k])
        for k in S2K:
            chk(f"S2_{n}:{k}", tw[f"s2_{n}"][k])
    for k in S1K:
        chk(f"S1_plan:{k}", tw["s1_plan"][k])
    for k in TK:
        chk(f"TOT:{k}", tw["total"][k])
    col, row = addr["SRC:total"]
    if abs(float(ws[f"{col}{row}"].value) - tw["s1_plan"]["reg"]) > 0.01:
        problems += 1; print("SRC total mismatch")
    if not quiet:
        t = tw["total"]
        print(f"итог (средние): выручка ${t['launch']:,.0f} | деньги ${t['net']:,.0f} (сразу ${t['now']:,.0f}, позже ${t['later']:,.0f}) "
              f"| расходы ${t['costs']:,.0f} | прибыль ${t['profit']:,.0f} | окупили: {t['ok']} | окупаемость вложений {t['payback']:.0%} "
              f"| трафика {t['roas']:.0%} | вложить до продаж ${t['start_cash']:,.0f}")
        for n in ("min", "mid", "max"):
            a, b = tw[f"s1_{n}"], tw[f"s2_{n}"]
            print(f"  {n}: этап 1 ${a['plan']:,.0f} (подписались {a['tg_sub_n']:.0f}) | этап 2 ${b['plan']:,.0f} | вместе ${a['plan'] + b['plan']:,.0f}")
    print("PROBLEMS:", problems)
    return problems

if __name__ == "__main__":
    sys.exit(1 if verify(sys.argv[1], sys.argv[2]) else 0)
