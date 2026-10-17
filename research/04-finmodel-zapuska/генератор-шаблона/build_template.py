#!/usr/bin/env python3
"""Заполняет шаблон заказчика «декомпозиция Узбекситан .xlsx» цифрами финмодели.

Листы по смыслу (согласовано): «через вебинармар» — холодный запуск, «анкета предзаписи» — тёплый запуск
в ТГ-канале, «гибридная» — шаги 1–4 шаблона + дожим ИИ-продавцом, деньги и рассрочки, расходы и прибыль,
касса по месяцам; «Трафик» — бюджет и цена регистрации. Трипваер не трогаем.

Использование: python3 build_template.py <out.xlsx> <addr.json>
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color

sys.path.insert(0, str(Path(__file__).parent))
from template_defaults import AI_SCEN, COLD, S3, S5, S6, TRAFFIC, WARM  # noqa: E402
from template_twin import cv_web_rate, warm_block, cold_block  # noqa: E402

TEMPLATE = "/Users/Bel-a/Downloads/декомпозиция Узбекситан .xlsx"
SA = "Модель продаж - анкета предзапи"
SB = "Модель продаж- через вебинармар"
SC = "Модель продаж - гибридная"
SD = "Трафик и лидмагнитная воронка"

THIN = Side(style="thin", color="FF000000")
NO = Side()
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

C_TITLE, C_NOTE, C_WHITE, C_TOT = "FF00FFFF", "FFFFF2CC", "FFFFFFFF", "FFB6D7A8"
C_SEC_A, C_SEC_B = "FFFFE599", "FFFFF2CC"
C_IN_A, C_IN_B = "FFCFE2F3", "FFC9DAF8"

FMT = {
    "int": "#,##0",
    "pct": "0.0%",
    "pct2": "0.00%",
    "usd": '"$"#,##0;[Red]-"$"#,##0',
    "usd2": '"$"#,##0.00',
    "roi": "0.00%",
    "mon": "0",
    "gen": "General",
}

ADDR = {}
TOKEN = re.compile(r"\{(?:(\w+):)?(\w+)\}")


def brd(l=False, r=False, t=False, b=False):  # noqa: E741
    return Border(left=THIN if l else NO, right=THIN if r else NO, top=THIN if t else NO, bottom=THIN if b else NO)


def mfont(size=11, bold=False, italic=False):
    return Font(name="Montserrat", size=size, bold=bold, italic=italic, color=Color(theme=1))


def solid(rgb):
    return PatternFill("solid", fgColor=rgb)


def col_shift(col, n):
    return chr(ord(col) + n)


def R(key, label, kind, fmt, v=None, note=None):
    """kind: in — вводная, f — формула, tot — итог (зелёный), tot_in — зелёная вводная."""
    return dict(key=key, label=label, kind=kind, fmt=fmt, v=v, note=note)


# ------------------------------------------------------------------ блоки
class Block:
    def __init__(self, bid, sheet, sections, lbl_col, val_col, first_row, sec_col=None, title=None, title_row=None,
                 title_fill=C_TITLE, title_size=11, sec_fill=C_SEC_A, in_fill=C_IN_A, title_height=None):
        self.bid, self.sheet, self.sections = bid, sheet, sections
        self.lbl_col, self.val_col, self.sec_col = lbl_col, val_col, sec_col
        self.first_row, self.title, self.title_row = first_row, title, title_row
        self.title_fill, self.title_size, self.sec_fill, self.in_fill = title_fill, title_size, sec_fill, in_fill
        self.title_height = title_height
        r = first_row
        for _label, rows in sections:
            for row in rows:
                row["r"] = r
                if (bid, row["key"]) in ADDR:
                    raise ValueError(f"duplicate {bid}:{row['key']}")
                ADDR[(bid, row["key"])] = (sheet, val_col, r)
                r += 1
        self.last_row = r - 1


def resolve(expr, bid, sheet):
    def rep(m):
        b = m.group(1) or bid
        sh, col, row = ADDR[(b, m.group(2))]
        return f"{col}{row}" if sh == sheet else f"'{sh}'!{col}{row}"
    return TOKEN.sub(rep, expr)


def merge_box(ws, c1, r1, c2, r2, value, font, fill, height=None):
    ws.merge_cells(f"{c1}{r1}:{c2}{r2}")
    ws[f"{c1}{r1}"] = value
    ws[f"{c1}{r1}"].font = font
    ws[f"{c1}{r1}"].fill = fill
    ws[f"{c1}{r1}"].alignment = CENTER
    for rr in range(r1, r2 + 1):
        for cc in range(ord(c1), ord(c2) + 1):
            ws[f"{chr(cc)}{rr}"].border = brd(l=cc == ord(c1), r=cc == ord(c2), t=rr == r1, b=rr == r2)
    if height:
        ws.row_dimensions[r1].height = height


def write_block(wb, blk):
    ws = wb[blk.sheet]
    if blk.title:
        left = blk.sec_col or blk.lbl_col
        merge_box(ws, left, blk.title_row, blk.val_col, blk.title_row, blk.title, mfont(blk.title_size, True),
                  solid(blk.title_fill), blk.title_height)
    for label, rows in blk.sections:
        r1, r2 = rows[0]["r"], rows[-1]["r"]
        if blk.sec_col and label:
            if r2 > r1:
                ws.merge_cells(f"{blk.sec_col}{r1}:{blk.sec_col}{r2}")
            c = ws[f"{blk.sec_col}{r1}"]
            c.value = label
            c.font = Font(name="Arial", bold=True, color=Color(theme=1))
            c.fill = solid(blk.sec_fill)
            c.alignment = CENTER
            for rr in range(r1, r2 + 1):
                ws[f"{blk.sec_col}{rr}"].border = brd(l=True, r=True, t=rr == r1, b=rr == r2)
        for row in rows:
            r, kind = row["r"], row["kind"]
            fill = {"in": blk.in_fill, "f": C_WHITE, "tot": C_TOT, "tot_in": C_TOT}[kind]
            bold_val = kind in ("tot", "tot_in")
            lc, vc = ws[f"{blk.lbl_col}{r}"], ws[f"{blk.val_col}{r}"]
            lc.value = row["label"]
            lc.font, lc.fill, lc.alignment, lc.border = mfont(11, True), solid(fill), CENTER, brd(l=True, r=True, b=True)
            v = row["v"]
            vc.value = resolve(v, blk.bid, blk.sheet) if isinstance(v, str) and v.startswith("=") else v
            vc.font, vc.fill, vc.alignment = mfont(11, bold_val), solid(fill), CENTER
            vc.border = brd(r=True, b=True)
            vc.number_format = FMT[row["fmt"]]
            if row.get("note"):
                cm = Comment(row["note"], "Claude")
                cm.width, cm.height = 320, 140
                vc.comment = cm


def clear_sheet(ws, rows=None):
    for mr in list(ws.merged_cells.ranges):
        if rows is None or (mr.min_row in rows):
            ws.unmerge_cells(str(mr))
    max_r = ws.max_row if rows is None else max(rows)
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max(ws.max_column, 16)):
        for c in row:
            if rows is not None and c.row not in rows:
                continue
            c.value = None
            c.font, c.fill, c.border, c.alignment = Font(), PatternFill(), Border(), Alignment()
            c.number_format = "General"
            c.comment = None
    for r in list(ws.row_dimensions.keys()):
        if rows is None or r in rows:
            ws.row_dimensions[r].height = None


def note_row(ws, label_cell, text_range, text, height):
    c1, c2 = text_range
    ws[label_cell] = "Суть модели:"
    ws[label_cell].font, ws[label_cell].fill, ws[label_cell].alignment = mfont(10, True), solid(C_NOTE), CENTER
    row = int(re.sub(r"\D", "", c1))
    ws.merge_cells(f"{c1}:{c2}")
    ws[c1] = text
    ws[c1].font, ws[c1].fill, ws[c1].alignment = mfont(10, True), solid(C_NOTE), CENTER
    ws.row_dimensions[row].height = height


# ------------------------------------------------------------------ содержание блоков
N = {
    "att": "Доходимость до вебинара: в среднем 15–20% (https://www.sostav.ru/blogs/287076/96331, 2026); 20–30% — норма, "
           "50% — хорошо (https://vc.ru/id972483/798734-dohodimost-do-vebinara-kakoi-procent-optimalnyi). База брифа рынка — 35%.",
    "conv": "Со вебинара покупают 5–20% зрителей (skillspace.ru, РФ). В брифе рынка: заявки — 12% участников, из них "
            "оплачивают 40%. Доли автооплат и предоплат отдельно не нашли — оценка.",
    "op_bk": "Заказ → оплата: 66,8% (https://getcourse.ru/magazine/analitika_2025, пик декабря 2025); 46% в примере "
             "финмодели школы (noboring-finance.ru).",
    "op_or": "Звонок → оплата 50–55% по гайду (не проверено); база брифа рынка — 40%.",
    "check": "Чек 300 $ — по ответу Васи. У конкурентов в УЗ: фриланс 170–595 $, ИИ-академии 255–510 $; базовый чек "
             "брифа — 212 $ (https://claude.ai/artifact/RGUpxwR5Sn9F54wb1Eq5Hb).",
    "peak": "Из шаблона: пик = 60% участников. Вынесено из формулы в ячейку.",
    "cv_web": "Считается из ставок ИИ-продавца и менеджера на листе «Модель продаж - гибридная», шаг 3: "
              "ответили × (оплатили сами + горячие × закрытие менеджера). База — только те, кто был и ничего не купил.",
    "cv_web_scen": "Дожим в боте тех, кто был на вебинаре и ничего не купил: ИИ-продавец + менеджер. Считается от "
                   "«участники − заказы − автооплаты − предоплаты», а не от всех участников.",
    "cv_dozhim": "Считается из ставок ИИ-продавца и менеджера на листе «Модель продаж - гибридная», шаг 3, с учётом того, "
                 "сколько людей дошли до продажи. База — зарегистрированные, кто ничего не купил.",
    "cv_dozhim_scen": "Дожим в боте ИИ-продавцом и менеджером. База — зарегистрированные на марафон, кто ничего не купил.",
    "churn": "Отписки 8–12% подписчиков (https://vc.ru/telegram/1425019-kak-poschitat-okupaemost-telegram-kanala, РФ, "
             "период не указан).",
    "warm_att": "Холодный марафон: день 1 — 25%, день 2 — 14% регистраций; тёплый: день 1 — 68%, на день 2 вернулись 72% "
                "(vc.ru, РФ).",
    "estimate": "Оценка — данных по Узбекистану нет.",
    "ai": "Оценка. Кейс вендора: конверсия в чате выросла с ~10–12% до 25% за месяц (sellergpt.ru, B2B) — для УЗ не проверено.",
    "acq": "Payme / Click / Uzum: 0,8–2%, подключают только местным юрлицам (https://uzneo.uz/blog/payme-click-uzum-podklyuchit-oplatu).",
    "partner": "Рассрочка 0% на 6–12 мес через Uzum и Alif — стандарт рынка УЗ (бриф). Alif одобряет 50–60% заявок "
               "(https://www.spot.uz/ru/2025/07/11/bnpl-regulation/), Uzum Nasiya — до 97%.",
    "fee": "Ставки Uzum и Alif для школ не найдены — запросить (задача CURSOR-q1z0). Ориентир РФ: 8,5% за 3 мес, 15% за 12 мес "
           "(https://prodamus.ru/chastyami).",
    "refund": "Возвраты 5% — допущение брифа рынка.",
    "ai_setup": "Разработка ИИ-бота в УЗ — 5–25 млн сум; ИИ и хостинг — 0,3–3 млн сум в месяц (https://innosoft.uz/blog/ai-narxi-2026).",
    "tax": "ООО — налог с оборота 4% до 1 млрд сум в год; ИП и самозанятые — 1%; выше порога — НДС 12% и налог на прибыль 15% "
           "(https://www.spot.uz/ru/2025/11/26/taxes-for-2026/).",
    "div": "Дивиденды нерезиденту — 10%, для России тоже 10% "
           "(https://taxsummaries.pwc.com/republic-of-uzbekistan/corporate/withholding-taxes).",
    "curator": "Кураторы 10–25 $ на ученика — допущение брифа рынка.",
    "m_pct": "Затраты на продажи 10% выручки — допущение брифа рынка.",
    "cpl": "Лид в образовании УЗ — 2,3–3,6 $ (https://aksanov.digital/blog/meta-ads/analiz-kampaniy); в сезон июль–сентябрь "
           "10–12 $ (https://t.me/s/targetologi_uz); база брифа — 6 $.",
    "tg_cpl": "Telegram-посевы: подписчик 0,59–1,02 $ (https://telegrad.uz/blog/skolko-stoit-reklama-v-telegram.html). "
              "В холодном запуске не используем — бюджет 0.",
}


def cold_sections(i, bid):
    base = i == 0
    ai = S3 if base else {k: v[i - 1] for k, v in AI_SCEN.items()}
    reg_v = "={B_src:total}" if base else COLD["reg"][i]
    traffic_v = "={D_tr:budget_total}" if base else COLD["traffic"][i]
    cv_web_v = "={C_s3:eng_att}*{C_s3:bot}*({C_s3:ai_pay}+{C_s3:ai_hot}*{C_s3:m_hot})" if base else cv_web_rate(ai)
    return [
        ("МАРКЕТИНГ И ВОРОНКА", [
            R("reg", "Регистраций:", "f" if base else "in", "int", reg_v,
              "Из таблицы «От куда?» справа: таргет — с листа «Трафик» (бюджет ÷ цена регистрации)." if base else None),
            R("att", "Доходимость:", "in", "pct", COLD["att"][i], N["att"] if base else None),
            R("part", "Участников на вебинаре:", "f", "int", "={reg}*{att}"),
            R("peak", "Пик:", "f", "int", "={part}*{peak_pct}"),
            R("peak_pct", "Пик, % от участников", "in", "pct", COLD["peak_pct"][i], N["peak"] if base else None),
            R("orders", "Оставили заказ:", "f", "int", "={part}*{cv_order}"),
            R("orders_sum", "Сумма заказов:", "f", "usd", "={orders}*{check}"),
            R("cv_order", "Конверсия из участника в заказ", "in", "pct", COLD["cv_order"][i], N["conv"] if base else None),
            R("autopay", "Автооплат шт", "f", "int", "={part}*{cv_autopay}"),
            R("autopay_sum", "Сумма автооплат", "f", "usd", "={autopay}*{check}"),
            R("cv_autopay", "Конверсия в автооплаты", "in", "pct", COLD["cv_autopay"][i]),
            R("prepay", "Предоплат шт", "f", "int", "={part}*{cv_prepay}"),
            R("cv_prepay", "Конверсия в предоплату", "in", "pct", COLD["cv_prepay"][i]),
        ]),
        ("ОТДЕЛ ПРОДАЖ", [
            R("op_leads", "Заявок на оп", "f", "int", "={orders}+{prepay}"),
            R("op_leads_sum", "Сумма заявок на оп", "f", "usd", "={op_leads}*{check}"),
            R("op_bk", "CV ОП - БРОНЬ", "in", "pct", COLD["op_bk"][i], N["op_bk"] if base else None),
            R("op_bk_plan", "План закрытые брони", "f", "int", "={op_bk}*{prepay}"),
            R("op_or", "СV ОП - ЗАКАЗ", "in", "pct", COLD["op_or"][i], N["op_or"] if base else None),
            R("op_or_plan", "План закрытые заявки", "f", "int", "={op_or}*{orders}"),
            R("cv_web", "CV ОП - ВЕБИНАРНЫЕ (дожим в боте тех, кто был и не купил)", "f" if base else "in", "pct2",
              cv_web_v, N["cv_web"] if base else N["cv_web_scen"]),
            R("web_plan", "План закрытия вебинарные", "f", "int", "={cv_web}*MAX(0,{part}-{orders}-{autopay}-{prepay})"),
        ]),
        ("ИТОГИ", [
            R("op_paid", "Оплат оп", "f", "int", "={web_plan}+{op_or_plan}+{op_bk_plan}"),
            R("op_sum", "Сумма с оп", "f", "usd", "={op_paid}*{check}"),
            R("traffic", "Вложили в трафик", "f" if base else "in", "usd", traffic_v,
              "С листа «Трафик»: сумма бюджетов." if base else None),
            R("roi", "Окупаемость трафика", "f", "roi", "=IF({traffic}>0,{plan}/{traffic},0)"),
            R("plan", "План продаж", "tot", "usd", "={op_sum}+{autopay_sum}"),
            R("check", "Ср. чек", "in", "usd", COLD["check"][i], N["check"] if base else None),
        ]),
    ]


def warm_sections(i, bid, cold_bid):
    base = i == 0
    if base:
        cv_v = ("=IF(MAX(0,{reg}-{autopay}-{bk}-{orders})>0,(MAX(0,{part}-{autopay}-{bk}-{orders})*{C_s3:eng_att}"
                "+MAX(0,{reg}-{part})*{C_s3:eng_no})*{C_s3:bot}*({C_s3:ai_pay}+{C_s3:ai_hot}*{C_s3:m_hot})"
                "/MAX(0,{reg}-{autopay}-{bk}-{orders}),0)")
    else:
        # статичная ставка сценария: считаем тем же способом по цифрам сценария
        cold_b = TWIN_COLD[cold_bid]
        ai = {k: v[i - 1] for k, v in AI_SCEN.items()}
        cv_v = warm_block(WARM, i, cold_b["reg"], cold_b["op_paid"] + cold_b["autopay"], ai=ai)["cv_dozhim"]
    return [
        ("ПРОГРЕВ", [
            R("ch_start", "Подписчиков в ТГ-канале на старте", "in", "int", WARM["ch_start"][i],
              "Если у эксперта уже есть канал — впиши. Это сильнее всего двигает тёплый запуск." if base else None),
            R("cold_reg", "Регистраций холодного запуска", "f", "int", "={%s:reg}" % cold_bid),
            R("ch_sub", "Подписались на канал, % регистраций", "in", "pct", WARM["ch_sub"][i], N["estimate"] if base else None),
            R("ch_from_cold", "Пришли в канал из холодного запуска", "f", "int", "={cold_reg}*{ch_sub}"),
            R("ch_org", "Органический прирост до тёплого запуска", "in", "int", WARM["ch_org"][i], N["estimate"] if base else None),
            R("ch_churn", "Отписки до тёплого запуска, %", "in", "pct", WARM["ch_churn"][i], N["churn"] if base else None),
            R("ch_buyers", "Уже купили курс в холодном — не целевые", "f", "int",
              "=({%s:op_paid}+{%s:autopay})*{ch_sub}" % (cold_bid, cold_bid)),
            R("base", "Средние охваты: база ТГ-канала к тёплому запуску", "f", "int",
              "=MAX(0,({ch_start}+{ch_from_cold}+{ch_org})*(1-{ch_churn})-{ch_buyers})"),
            R("reg", "Людей в анкете: регистраций на марафон", "f", "int", "={base}*{cv_reg}"),
            R("cv_reg", "CV охват - анкета: база канала → регистрация", "in", "pct", WARM["cv_reg"][i], N["estimate"] if base else None),
        ]),
        ("МАРКЕТИНГ ВОРОНКА", [
            R("part", "Дошли до продающего дня марафона", "f", "int", "={reg}*{att1}*{att2}"),
            R("att1", "Доходимость: день 1, % регистраций", "in", "pct", WARM["att1"][i], N["warm_att"] if base else None),
            R("att2", "Доходимость: день 2 (продажа), % от дня 1", "in", "pct", WARM["att2"][i]),
            R("autopay", "Автооплаты", "f", "int", "={part}*{cv_autopay}"),
            R("cv_autopay", "СV автооплата - участники", "in", "pct", WARM["cv_autopay"][i]),
            R("bk", "Брони", "f", "int", "={part}*{cv_bk}"),
            R("cv_bk", "СV бронь - участники", "in", "pct", WARM["cv_bk"][i]),
            R("orders", "Заказы", "f", "int", "={part}*{cv_order}"),
            R("cv_order", "CV заказ - участники", "in", "pct", WARM["cv_order"][i]),
            R("op_leads", "Заявок в работу на ОП", "f", "int", "={bk}+{orders}"),
        ]),
        ("ОТДЕЛ ПРОДАЖ", [
            R("op_bk", "CV ОП - БРОНЬ", "in", "pct", WARM["op_bk"][i], N["op_bk"] if base else None),
            R("op_bk_plan", "План закрытые брони", "f", "int", "={bk}*{op_bk}"),
            R("op_or", "СV ОП - ЗАКАЗ", "in", "pct", WARM["op_or"][i], N["op_or"] if base else None),
            R("op_or_plan", "План закрытые заявки", "f", "int", "={orders}*{op_or}"),
            R("cv_dozhim", "CV ОП - АНКЕТА: дожим в боте (ИИ-продавец + менеджер)", "f" if base else "in", "pct2", cv_v,
              N["cv_dozhim"] if base else N["cv_dozhim_scen"]),
            R("dozhim_plan", "План закрытия анкеты: дожим тех, кто не купил", "f", "int",
              "=MAX(0,{reg}-{autopay}-{bk}-{orders})*{cv_dozhim}"),
        ]),
        ("ИТОГИ", [
            R("op_total", "Итого продаж ОП", "tot", "int", "={dozhim_plan}+{op_or_plan}+{op_bk_plan}"),
            R("op_sum", "План ОП сумма", "tot", "usd", "={op_total}*{check}"),
            R("autopay_total", "Итого автооплат", "tot", "int", "={autopay}"),
            R("total", "Итог общ", "tot", "int", "={op_total}+{autopay_total}"),
            R("check", "Средний чек", "tot_in", "usd", WARM["check"][i], N["check"] if base else None),
            R("plan", "План продаж", "tot", "usd", "={total}*{check}"),
        ]),
    ]


TWIN_COLD = {}


def build(out_path, addr_path):
    wb = load_workbook(TEMPLATE)

    # --- сценарные цифры холодного запуска для статичных ставок тёплого
    for j, bid in enumerate(["B_min", "B_mid", "B_max"], start=1):
        ai = {k: v[j - 1] for k, v in AI_SCEN.items()}
        TWIN_COLD[bid] = cold_block(COLD, j, COLD["reg"][j], COLD["traffic"][j], cv_web_rate(ai))

    # ================= раскладка всех блоков (адреса нужны до записи формул)
    # Трафик: таблица 1
    ADDR[("D_tr", "target_budget")] = (SD, "C", 6)
    ADDR[("D_tr", "target_reg")] = (SD, "D", 6)
    ADDR[("D_tr", "budget_total")] = (SD, "C", 9)
    ADDR[("D_tr", "reg_total")] = (SD, "D", 9)
    # «От куда?»
    ADDR[("B_src", "total")] = (SB, "I", 8)
    ADDR[("B_src", "target")] = (SB, "I", 9)

    blocks = []
    # --- лист «через вебинармар»: основной блок и сценарии
    b_base = Block("B_base", SB, cold_sections(0, "B_base"), "C", "D", 6, sec_col="B",
                   title="План декомпозиции продаж через вебинар — холодный запуск с таргета", title_row=5,
                   title_size=12, sec_fill=C_SEC_B, in_fill=C_IN_B, title_height=36)
    blocks.append(b_base)
    b_hdr = b_base.last_row + 3
    b_title = b_hdr + 4
    b_scen = []
    for j, (bid, cols, title, fill) in enumerate([
            ("B_min", ("B", "C", "D"), "ПЛАН МИНИМУМ", "FFF4CCCC"),
            ("B_mid", ("G", "H", "I"), "ПЛАН СРЕДНИЙ", "FFFCE5CD"),
            ("B_max", ("L", "M", "N"), "ПЛАН МАКСИМУМ", "FFB6D7A8")], start=1):
        blk = Block(bid, SB, cold_sections(j, bid), cols[1], cols[2], b_title + 1, sec_col=cols[0], title=title,
                    title_row=b_title, title_fill=fill, sec_fill=C_SEC_B, in_fill=C_IN_B)
        b_scen.append(blk)
        blocks.append(blk)

    # --- лист «анкета предзаписи»: тёплый запуск
    a_base = Block("A_base", SA, warm_sections(0, "A_base", "B_base"), "C", "D", 7, sec_col="B",
                   title="Декомпозиция тёплого запуска в ТГ-канале — марафон 2 дня", title_row=6, title_height=36)
    blocks.append(a_base)
    a_hdr = a_base.last_row + 4
    a_title = a_hdr + 5
    for j, (bid, cold_bid, cols, title, fill) in enumerate([
            ("A_min", "B_min", ("B", "C", "D"), "ПЛАН МИНИМУМ", "FFF4CCCC"),
            ("A_opt", "B_mid", ("G", "H", "I"), "ПЛАН ОПТИМУМ", "FFFFF2CC"),
            ("A_max", "B_max", ("L", "M", "N"), "ПЛАН МАКСИМУМ", "FF00FF00")], start=1):
        blocks.append(Block(bid, SA, warm_sections(j, bid, cold_bid), cols[1], cols[2], a_title + 1, sec_col=cols[0],
                            title=title, title_row=a_title, title_fill=fill))

    # --- лист «гибридная»
    s1_sections = []
    for label, rows in warm_sections(0, "A_base_mirror", "B_base"):
        mirrored = [R(r["key"], r["label"], "tot" if r["kind"] in ("tot", "tot_in") else "f", r["fmt"],
                      "={A_base:%s}" % r["key"]) for r in rows]
        s1_sections.append((label, mirrored))
    s1 = Block("C_s1", SC, s1_sections, "C", "D", 6, sec_col="B",
               title="ШАГ 1. ТЁПЛЫЙ ЗАПУСК В ТГ-КАНАЛЕ (цифры — с листа «анкета предзаписи»)", title_row=5, title_height=48)
    s2 = Block("C_s2", SC, [(None, [
        R("reg", "Регистраций:", "f", "int", "={B_base:reg}"),
        R("att", "Доходимость:", "f", "pct", "={B_base:att}"),
        R("part", "Участников на вебинаре:", "f", "int", "={B_base:part}"),
        R("peak", "Пик:", "f", "int", "={B_base:peak}"),
        R("orders", "Оставили заказ:", "f", "int", "={B_base:orders}"),
        R("orders_sum", "Сумма заказов:", "f", "usd", "={B_base:orders_sum}"),
        R("cv_order", "Конверсия из участника в заказ", "f", "pct", "={B_base:cv_order}"),
        R("autopay", "Автооплат шт", "f", "int", "={B_base:autopay}"),
        R("autopay_sum", "Сумма автооплат", "f", "usd", "={B_base:autopay_sum}"),
        R("cv_autopay", "Конверсия в автооплаты", "f", "pct", "={B_base:cv_autopay}"),
        R("prepay", "Предоплат шт", "f", "int", "={B_base:prepay}"),
        R("cv_prepay", "Конверсия в предоплату", "f", "pct", "={B_base:cv_prepay}"),
        R("op_leads", "Заявок на оп", "f", "int", "={B_base:op_leads}"),
        R("op_leads_sum", "Сумма заявок на оп", "f", "usd", "={B_base:op_leads_sum}"),
        R("op_paid", "Оплат оп (брони + заказы)", "f", "int", "={B_base:op_bk_plan}+{B_base:op_or_plan}",
          "Без дожима — дожим считаем в шаге 3."),
        R("cv_op", "Конверсия ОП", "f", "pct", "=IF({op_leads}>0,{op_paid}/{op_leads},0)",
          "Средняя по броням и заказам: ставки отдельно — на листе «через вебинармар»."),
        R("op_sum", "Сумма с оп", "f", "usd", "={op_paid}*{check}"),
        R("traffic", "Вложили в трафик", "f", "usd", "={B_base:traffic}"),
        R("roi", "Окупаемость трафика", "f", "roi", "=IF({traffic}>0,{plan}/{traffic},0)"),
        R("plan", "План продаж", "tot", "usd", "={op_sum}+{autopay_sum}"),
        R("check", "Ср. чек", "f", "usd", "={B_base:check}"),
    ])], "G", "H", 6, title="ШАГ 2. ХОЛОДНЫЙ ЗАПУСК: ТАРГЕТ → ВЕБИНАР (цифры — с листа «через вебинармар»)",
        title_row=5, in_fill=C_IN_B)
    s3 = Block("C_s3", SC, [(None, [
        R("reg", "Регистраций:", "f", "int", "={C_s2:reg}"),
        R("noact", "Были на вебинаре, но не купили", "f", "int", "=MAX(0,{C_s2:part}-{C_s2:orders}-{C_s2:autopay}-{C_s2:prepay})"),
        R("noshow", "Не дошли до вебинара", "f", "int", "=MAX(0,{C_s2:reg}-{C_s2:part})"),
        R("bot", "В боте, % регистраций", "in", "pct", S3["bot"], "Регистрация идёт через Telegram-бота — 100%."),
        R("eng_att", "ИИ-продавец: ответили — были на вебинаре, %", "in", "pct", S3["eng_att"], N["ai"]),
        R("eng_no", "ИИ-продавец: ответили — не дошли, %", "in", "pct", S3["eng_no"], N["ai"]),
        R("dialogs", "Диалогов ИИ-продавца", "f", "int", "=({noact}*{eng_att}+{noshow}*{eng_no})*{bot}"),
        R("ai_pay", "ИИ-продавец: оплатили сами по ссылке, % диалогов", "in", "pct", S3["ai_pay"], N["ai"]),
        R("ai_paid", "Оплат через ИИ-продавца", "f", "int", "={dialogs}*{ai_pay}"),
        R("ai_hot", "CV в заявку: горячий лид менеджеру, % диалогов", "in", "pct", S3["ai_hot"], N["ai"]),
        R("hot", "Заявок: горячих лидов менеджеру", "f", "int", "={dialogs}*{ai_hot}"),
        R("m_hot", "CV менеджер: горячий лид → оплата", "in", "pct", S3["m_hot"], "Оценка: лид уже прогрет ИИ-продавцом."),
        R("m_paid", "Оплат менеджера с горячих лидов", "f", "int", "={hot}*{m_hot}"),
        R("total", "Итого продаж дожима", "f", "int", "={ai_paid}+{m_paid}"),
        R("plan", "План продаж", "tot", "usd", "={total}*{C_s2:check}"),
    ])], "K", "L", 6, title="ШАГ 3. ДОЖИМ В БОТЕ: ИИ-ПРОДАВЕЦ + МЕНЕДЖЕР", title_row=5, in_fill=C_IN_B)
    s4_title = s3.last_row + 2
    s4 = Block("C_s4", SC, [(None, [
        R("op_plan", "ПЛАН ОП ПРОДАЖ", "f", "usd", "={C_s1:op_sum}+{C_s2:op_sum}+{C_s3:plan}"),
        R("launch_plan", "ПЛАН ПРОДАЖ ЗАПУСК", "tot", "usd", "={C_s1:plan}+{C_s2:plan}+{C_s3:plan}"),
    ])], "K", "L", s4_title + 1, title="ШАГ 4. ОБЩИЙ ИТОГ", title_row=s4_title)
    band = max(s1.last_row, s2.last_row, s4.last_row) + 3
    s5 = Block("C_s5", SC, [
        ("ОПЛАТЫ", [
            R("full", "Полная оплата, % продаж", "in", "pct", S5["full"]),
            R("partner", "Рассрочка Uzum Nasiya / Alif, % продаж", "in", "pct", S5["partner"], N["partner"]),
            R("own", "Своя рассрочка школы, % продаж (остаток)", "f", "pct", "=MAX(0,1-{full}-{partner})"),
            R("acq", "Эквайринг Payme / Click / Uzum, %", "in", "pct", S5["acq"], N["acq"]),
            R("fee", "Комиссия рассрочки Uzum / Alif для школы, %", "in", "pct", S5["fee"], N["fee"]),
            R("term", "Своя рассрочка: срок, мес (1–12)", "in", "mon", S5["term"]),
            R("first", "Своя рассрочка: первый платёж, % цены", "in", "pct", S5["first"], N["estimate"]),
            R("default", "Своя рассрочка: не доплатят, % остатка", "in", "pct", S5["default"], N["estimate"]),
            R("refund", "Возвраты, % выручки", "in", "pct", S5["refund"], N["refund"]),
            R("gap", "Тёплый запуск через, мес после холодного (1–6)", "in", "mon", S5["gap"]),
        ]),
        ("ДЕНЬГИ", [
            R("cold_rev", "Выручка холодного запуска (шаг 2 + шаг 3)", "f", "usd", "={C_s2:plan}+{C_s3:plan}"),
            R("warm_rev", "Выручка тёплого запуска (шаг 1)", "f", "usd", "={C_s1:plan}"),
            R("rev", "Выручка цикла", "tot", "usd", "={cold_rev}+{warm_rev}"),
            R("term_n", "Срок своей рассрочки в расчёте, мес", "f", "mon", "=MAX(1,MIN(12,ROUND({term},0)))"),
            R("first_eff", "Первый платёж в расчёте, %", "f", "pct", "=IF({term_n}<=1,1,{first})"),
            R("cold_now", "Холодный: пришло сразу (после возвратов)", "f", "usd",
              "={cold_rev}*({full}*(1-{acq})+{partner}*(1-{fee})+{own}*{first_eff}*(1-{acq}))-{cold_rev}*{refund}"),
            R("cold_month", "Холодный: своя рассрочка, $ в месяц со 2-го месяца", "f", "usd",
              "=IF({term_n}>1,{cold_rev}*{own}*(1-{first_eff})*(1-{default})*(1-{acq})/({term_n}-1),0)"),
            R("warm_now", "Тёплый: пришло сразу (после возвратов)", "f", "usd",
              "={warm_rev}*({full}*(1-{acq})+{partner}*(1-{fee})+{own}*{first_eff}*(1-{acq}))-{warm_rev}*{refund}"),
            R("warm_month", "Тёплый: своя рассрочка, $ в месяц после запуска", "f", "usd",
              "=IF({term_n}>1,{warm_rev}*{own}*(1-{first_eff})*(1-{default})*(1-{acq})/({term_n}-1),0)"),
            R("now", "Пришло сразу, в месяц продаж", "f", "usd", "={cold_now}+{warm_now}"),
            R("later", "Придёт позже по своей рассрочке", "f", "usd", "=({cold_month}+{warm_month})*({term_n}-1)"),
            R("loss_acq", "Потери: эквайринг", "f", "usd",
              "={rev}*({full}+{own}*{first_eff}+{own}*(1-{first_eff})*(1-{default}))*{acq}"),
            R("loss_fee", "Потери: комиссия рассрочки Uzum / Alif", "f", "usd", "={rev}*{partner}*{fee}"),
            R("loss_def", "Потери: не доплатили по своей рассрочке", "f", "usd", "={rev}*{own}*(1-{first_eff})*{default}"),
            R("loss_ref", "Потери: возвраты", "f", "usd", "={rev}*{refund}"),
            R("net", "Чистые деньги цикла", "tot", "usd", "={now}+{later}"),
        ]),
    ], "C", "D", band + 1, sec_col="B", title="ШАГ 5. ДЕНЬГИ И РАССРОЧКИ", title_row=band)
    s6 = Block("C_s6", SC, [(None, [
        R("ads", "Реклама: таргет (с листа «Трафик»)", "f", "usd", "={D_tr:budget_total}"),
        R("ai_setup", "ИИ-продавец: разработка, разово", "in", "usd", S6["ai_setup"], N["ai_setup"]),
        R("ai_sub", "ИИ-продавец: ИИ и хостинг, $ в месяц", "in", "usd", S6["ai_sub"], N["ai_setup"]),
        R("ai_dialog", "ИИ-продавец: $ за 1 диалог", "in", "usd2", S6["ai_dialog"], N["estimate"]),
        R("salary", "Менеджер: оклад, $ в месяц", "in", "usd", S6["salary"], N["estimate"]),
        R("m_pct", "Менеджер: % с продаж ОП", "in", "pct", S6["m_pct"], N["m_pct"]),
        R("team", "Команда запуска на цикл", "in", "usd", S6["team"], "Таргетолог, техспец, дизайнер, монтаж. " + N["estimate"]),
        R("services", "Сервисы: бот, стрим, лендинг, $ в месяц", "in", "usd", S6["services"], N["estimate"]),
        R("curator", "Кураторы на 1 ученика", "in", "usd", S6["curator"], N["curator"]),
        R("tax", "Налог с оборота, %", "in", "pct", S6["tax"], N["tax"]),
        R("months", "Месяцев в цикле", "f", "mon", "={C_s5:gap}+1"),
        R("cold_students", "Учеников холодного запуска", "f", "int", "={C_s2:op_paid}+{C_s2:autopay}+{C_s3:total}"),
        R("warm_students", "Учеников тёплого запуска", "f", "int", "={C_s1:total}"),
        R("cold_mgr_rev", "Выручка менеджера: холодный", "f", "usd", "=({C_s2:op_paid}+{C_s3:m_paid})*{C_s2:check}"),
        R("warm_mgr_rev", "Выручка менеджера: тёплый", "f", "usd", "={C_s1:op_total}*{C_s1:check}"),
        R("cost_ai", "Расходы: ИИ-продавец", "f", "usd", "={ai_setup}+{ai_sub}*{months}+{C_s3:dialogs}*{ai_dialog}"),
        R("cost_mgr", "Расходы: менеджер", "f", "usd", "={salary}*{months}+({cold_mgr_rev}+{warm_mgr_rev})*{m_pct}"),
        R("cost_team", "Расходы: команда", "f", "usd", "={team}"),
        R("cost_services", "Расходы: сервисы", "f", "usd", "={services}*{months}"),
        R("cost_cur", "Расходы: кураторы", "f", "usd", "=({cold_students}+{warm_students})*{curator}"),
        R("cost_tax", "Расходы: налог", "f", "usd", "={C_s5:net}*{tax}"),
        R("costs", "РАСХОДЫ ВСЕГО", "tot", "usd",
          "={ads}+{cost_ai}+{cost_mgr}+{cost_team}+{cost_services}+{cost_cur}+{cost_tax}"),
        R("net", "Чистые деньги цикла (шаг 5)", "f", "usd", "={C_s5:net}"),
        R("profit", "ПРИБЫЛЬ ЦИКЛА", "tot", "usd", "={net}-{costs}"),
        R("profit_rec", "Прибыль без разовой разработки ИИ", "f", "usd", "={profit}+{ai_setup}"),
        R("partner_share", "Доля эксперта / партнёра, %", "in", "pct", S6["partner_share"]),
        R("div_tax", "Налог на дивиденды, %", "in", "pct", S6["div_tax"], N["div"]),
        R("owner", "Собственнику после доли и налога на дивиденды", "tot", "usd",
          "=({profit}-MAX(0,{profit})*{partner_share})-MAX(0,{profit}-MAX(0,{profit})*{partner_share})*{div_tax}"),
        R("romi", "ROMI: прибыль ÷ реклама", "f", "roi", "=IF({ads}>0,{profit}/{ads},0)"),
        R("start_cash", "Нужно вложить до первых продаж", "f", "usd", "={ads}+{ai_setup}+{team}+{ai_sub}+{salary}+{services}"),
        R("cold_var", "Для кассы: расходы в месяц холодного запуска", "f", "usd",
          "={ads}+{ai_setup}+{C_s3:dialogs}*{ai_dialog}+{cold_mgr_rev}*{m_pct}+{cold_students}*{curator}"),
        R("warm_var", "Для кассы: расходы в месяц тёплого запуска", "f", "usd", "={warm_mgr_rev}*{m_pct}+{warm_students}*{curator}"),
        R("fixed_month", "Для кассы: постоянные расходы в месяц", "f", "usd", "={ai_sub}+{salary}+{services}+{team}/{months}"),
    ])], "G", "H", band + 1, title="ШАГ 6. РАСХОДЫ И ПРИБЫЛЬ", title_row=band)
    blocks += [s1, s2, s3, s4, s5, s6]

    # касса по месяцам: адреса
    cash_top = band
    cash_first = band + 2
    cash_months = 18
    cash_total = cash_first + cash_months
    ADDR[("C_cash", "low")] = (SC, "O", cash_total + 4)

    # ================= очистка листов и запись
    for name in (SA, SB, SC):
        clear_sheet(wb[name])
    clear_sheet(wb[SD], rows=set(range(1, 12)))

    for blk in blocks:
        write_block(wb, blk)

    # --- «анкета»: суть модели, заголовок сценариев, ширины
    ws = wb[SA]
    note_row(ws, "B3", ("C3", "H3"),
             "Холодный запуск с таргета приводит людей в Telegram-канал. Раз в 2 месяца в канале проводим тёплый запуск — "
             "двухдневный марафон. На марафоне автооплаты, брони и заказы; брони и заказы закрывает менеджер, "
             "остальных дожимает ИИ-продавец в боте.", 66)
    merge_box(ws, "B", a_hdr, "N", a_hdr + 1, "НИЖЕ РАЗНЫЕ СЦЕНАРИИ ПЛАНА ПРОДАЖ", mfont(13, True), solid(C_NOTE))
    for rr in (a_hdr, a_hdr + 1):
        for cc in "BCDEFGHIJKLMN":
            ws[f"{cc}{rr}"].border = Border()
    for col, w in {"B": 13, "C": 40, "D": 22, "G": 13, "H": 40, "I": 18, "L": 13, "M": 40, "N": 18}.items():
        ws.column_dimensions[col].width = w

    # --- «через вебинармар»: суть модели, «От куда?», сценарии, ширины
    ws = wb[SB]
    note_row(ws, "C2", ("D2", "J2"),
             "Таргет в Instagram и Facebook ведёт на регистрацию в Telegram-бота, дальше бесплатный вебинар. На вебинаре "
             "автооплаты, предоплаты и заказы; предоплаты и заказы закрывает менеджер, тех, кто был и не купил, "
             "дожимает ИИ-продавец в боте.", 66)
    merge_box(ws, "H", 5, "J", 5, "Декомпозиция регистрации", mfont(11, True), solid(C_TITLE), 36)
    for col, text, fill in (("H", "От куда?", "FFFFE599"), ("I", "План", "FFFFE599"), ("J", "Факт", "FFD9EAD3")):
        merge_box(ws, col, 6, col, 7, text, mfont(11, True), solid(fill))
    sources = [("Итого", "=SUM(I9:I19)", True), ("Таргет (Instagram / Facebook)", "='%s'!D6" % SD, False),
               ("Телеграм канал", 0, False), ("Бот", 0, False), ("Емайл", 0, False), ("Спам тг", 0, False),
               ("Спам вотс ап", 0, False), ("СМС", 0, False), ("ВКонтакте", 0, False),
               ("Каналы прошлых мероприятий", 0, False), ("Отдел продаж", 0, False), ("Блогеры", 0, False)]
    for k, (label, val, bold) in enumerate(sources):
        r = 8 + k
        for col, v in (("H", label), ("I", val), ("J", None)):
            c = ws[f"{col}{r}"]
            c.value = v
            c.font = mfont(11, bold or col == "H" and k == 1)
            c.alignment = CENTER
            c.border = brd(l=col == "H", r=True, b=True)
            c.number_format = FMT["int"]
    ws["I9"].comment = Comment("Регистрации с таргета: лист «Трафик», бюджет ÷ цена регистрации.", "Claude")
    ws["H10"].comment = Comment("Новая школа: своей базы на старте нет — 0. Если есть канал, бот или блогеры — впиши план.", "Claude")
    merge_box(ws, "B", b_hdr, "N", b_hdr + 1, "НИЖЕ РАЗНЫЕ СЦЕНАРИИ ПЛАНА ПРОДАЖ", mfont(13, True), solid(C_NOTE))
    for rr in (b_hdr, b_hdr + 1):
        for cc in "BCDEFGHIJKLMN":
            ws[f"{cc}{rr}"].border = Border()
    for col, w in {"B": 14, "C": 40, "D": 22, "G": 14, "H": 40, "I": 20, "J": 14, "L": 14, "M": 40, "N": 18}.items():
        ws.column_dimensions[col].width = w

    # --- «гибридная»: суть модели, касса по месяцам, ширины
    ws = wb[SC]
    note_row(ws, "B3", ("C3", "O3"),
             "Один цикл. Холодный запуск: таргет → регистрация в бота → вебинар → продажи. Дальше дожим в боте: "
             "ИИ-продавец пишет текстом, горячих закрывает менеджер. Люди подписываются на Telegram-канал, через 2 месяца "
             "в канале тёплый запуск — марафон 2 дня. Ниже — деньги с рассрочками, расходы, прибыль и касса по месяцам.", 58)
    merge_box(ws, "K", cash_top, "O", cash_top, "ДЕНЬГИ ПО МЕСЯЦАМ", mfont(11, True), solid(C_TITLE))
    heads = ["Месяц", "Поступления", "Расходы", "Итог месяца", "Нарастающим итогом"]
    for col, h in zip("KLMNO", heads):
        c = ws[f"{col}{cash_top + 1}"]
        c.value, c.font, c.fill, c.alignment = h, mfont(11, True), solid(C_SEC_A), CENTER
        c.border = brd(l=True, r=True, b=True)
    ref = lambda b, k: resolve("{%s:%s}" % (b, k), "C_cash", SC)  # noqa: E731
    for n in range(cash_months):
        r = cash_first + n
        m = f"$K{r}"
        ws[f"K{r}"] = n + 1
        ws[f"L{r}"] = (f"=IF({m}=1,{ref('C_s5','cold_now')},0)"
                       f"+IF(AND({m}>=2,{m}<={ref('C_s5','term_n')}),{ref('C_s5','cold_month')},0)"
                       f"+IF({m}={ref('C_s6','months')},{ref('C_s5','warm_now')},0)"
                       f"+IF(AND({m}>={ref('C_s6','months')}+1,{m}<={ref('C_s6','months')}+{ref('C_s5','term_n')}-1),"
                       f"{ref('C_s5','warm_month')},0)")
        ws[f"M{r}"] = (f"=IF({m}=1,{ref('C_s6','cold_var')},0)+IF({m}={ref('C_s6','months')},{ref('C_s6','warm_var')},0)"
                       f"+IF({m}<={ref('C_s6','months')},{ref('C_s6','fixed_month')},0)+{ref('C_s6','tax')}*L{r}")
        ws[f"N{r}"] = f"=L{r}-M{r}"
        ws[f"O{r}"] = f"=N{r}" if n == 0 else f"=O{r - 1}+N{r}"
        for col in "KLMNO":
            c = ws[f"{col}{r}"]
            c.font, c.alignment, c.border = mfont(11, col == "O"), CENTER, brd(l=col == "K", r=True, b=True)
            c.number_format = FMT["mon"] if col == "K" else FMT["usd"]
    last = cash_first + cash_months - 1
    rows_tail = [
        (cash_total, "Итого", {"L": f"=SUM(L{cash_first}:L{last})", "M": f"=SUM(M{cash_first}:M{last})",
                               "N": f"=SUM(N{cash_first}:N{last})"}, C_TOT),
        (cash_total + 2, "Сверка поступлений с шагом 5 (должно быть 0)",
         {"L": f"=ROUND(L{cash_total}-{ref('C_s5','net')},2)"}, None),
        (cash_total + 3, "Сверка расходов с шагом 6 (должно быть 0)",
         {"M": f"=ROUND(M{cash_total}-{ref('C_s6','costs')},2)"}, None),
        (cash_total + 4, "Самая низкая точка кассы", {"O": f"=MIN(0,MIN(O{cash_first}:O{last}))"}, None),
    ]
    for r, label, cells, fill in rows_tail:
        c = ws[f"K{r}"]
        c.value, c.font, c.alignment = label, mfont(11, True), CENTER
        if fill:
            c.fill = solid(fill)
        c.border = brd(l=True, r=True, b=True)
        for col in "LMNO":
            cc = ws[f"{col}{r}"]
            cc.value = cells.get(col)
            cc.font, cc.alignment, cc.number_format = mfont(11, True), CENTER, FMT["usd"]
            cc.border = brd(r=True, b=True)
            if fill:
                cc.fill = solid(fill)
        ws.row_dimensions[r].height = 32
    for col, w in {"B": 14, "C": 42, "D": 20, "G": 42, "H": 20, "K": 42, "L": 20, "M": 16, "N": 16, "O": 20}.items():
        ws.column_dimensions[col].width = w

    # --- «Трафик»: таблица 1
    ws = wb[SD]
    merge_box(ws, "B", 4, "H", 4, "Декомпозиция трафика — регистрации на холодный запуск (вебинар)", mfont(11, True),
              solid(C_TITLE), 30)
    heads = ["Площадки", "Бюджет", "Подписчики (регистрации в бота)", "ЦП (цена регистрации)",
             "Активации (дошли до вебинара)", "Предзапись (заказ, предоплата или автооплата)",
             "Стоимость активации (цена участника)"]
    for col, h in zip("BCDEFGH", heads):
        c = ws[f"{col}5"]
        c.value, c.font, c.alignment, c.border = h, mfont(11), CENTER, brd(l=True, r=True, t=True, b=True)
    ws.row_dimensions[5].height = 62
    b = lambda k: resolve("{B_base:%s}" % k, "D_tr", SD)  # noqa: E731
    platforms = [(6, "Таргет Instagram / Facebook", TRAFFIC["target_budget"], TRAFFIC["target_cpl"], N["cpl"]),
                 (7, "Тг канал", TRAFFIC["tg_budget"], TRAFFIC["tg_cpl"], N["tg_cpl"]),
                 (8, "ВК", TRAFFIC["vk_budget"], TRAFFIC["vk_cpl"], "В Узбекистане ВК не используем — 0.")]
    for r, label, budget, cpl, note in platforms:
        vals = {"B": label, "C": budget, "D": f"=IF(E{r}>0,C{r}/E{r},0)", "E": cpl, "F": f"=D{r}*{b('att')}",
                "G": f"=F{r}*({b('cv_order')}+{b('cv_autopay')}+{b('cv_prepay')})", "H": f"=IF(F{r}>0,C{r}/F{r},0)"}
        fmts = {"B": "gen", "C": "usd", "D": "int", "E": "usd2", "F": "int", "G": "int", "H": "usd2"}
        for col, v in vals.items():
            c = ws[f"{col}{r}"]
            c.value, c.font, c.alignment = v, mfont(11, col == "B"), CENTER
            c.border = brd(l=col == "B", r=True, b=True)
            c.number_format = FMT[fmts[col]]
            if col in "CE":
                c.fill = solid(C_IN_A)
        cm = Comment(note, "Claude")
        cm.width, cm.height = 320, 120
        ws[f"E{r}"].comment = cm
    tot = {"B": "Итог", "C": "=SUM(C6:C8)", "D": "=SUM(D6:D8)", "E": "=IF(D9>0,C9/D9,0)", "F": "=SUM(F6:F8)",
           "G": "=SUM(G6:G8)", "H": "=IF(F9>0,C9/F9,0)"}
    fmts = {"B": "gen", "C": "usd", "D": "int", "E": "usd2", "F": "int", "G": "int", "H": "usd2"}
    for col, v in tot.items():
        c = ws[f"{col}9"]
        c.value, c.font, c.fill, c.alignment = v, mfont(11, True), solid("FF00FF00"), CENTER
        c.border = brd(l=col == "B", r=True, b=True)
        c.number_format = FMT[fmts[col]]
    ws["E9"].comment = Comment("Средняя цена регистрации = бюджет ÷ регистрации (в шаблоне было простое среднее цен).", "Claude")
    ws["B11"] = ("Трипваер ниже в этой воронке не используется — блок оставлен из шаблона без изменений (цифры в рублях).")
    ws["B11"].font = mfont(9, italic=True)
    for col, w in {"B": 26, "C": 14, "D": 18, "E": 16, "F": 18, "G": 22, "H": 20}.items():
        ws.column_dimensions[col].width = max(w, ws.column_dimensions[col].width or 0)

    wb.save(out_path)
    Path(addr_path).write_text(json.dumps({f"{k[0]}:{k[1]}": v for k, v in ADDR.items()}, ensure_ascii=False, indent=0))
    print("ok", out_path, "cells:", len(ADDR), "| band", band, "| B scen title row", b_title, "| A scen title row", a_title)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
