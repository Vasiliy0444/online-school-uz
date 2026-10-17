#!/usr/bin/env python3
"""Собирает «Финмодель запуска УЗ.xlsx»: формулы, вводные, 3 сценария, касса по месяцам.

Использование:
  python3 build_model.py <out.xlsx> rowmap.json   — собрать книгу (вводные — model_defaults.py)
  python3 <скилл xlsx>/scripts/recalc.py <out.xlsx> 90   — пересчитать формулы в LibreOffice
  python3 verify_model.py <out.xlsx> rowmap.json   — сверить с Python-двойником (twin.py)
  python3 stress_test.py   — стресс-тесты (путь к recalc.py внутри файла)
Внимание: пересборка перезапишет правки, сделанные руками в самой книге.
"""
import json
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from content_sources import BRIEF_URL, SOURCE_ALIAS, SOURCES  # noqa: E402
from content_texts import FINDINGS, SHEETS_OVERVIEW  # noqa: E402
from model_defaults import DEF  # noqa: E402

S1, S2, S3, S4 = "Декомпозиция цикла", "Деньги по месяцам", "Разбор исходника", "Источники цифр"
SCEN = ["Минимум", "Оптимум", "Максимум"]
COLS = ["D", "E", "F"]
FIRST_ROW = 7

# Лист 2: по 4 колонки на сценарий — поступления, расходы, итог месяца, нарастающим итогом
S2_BLOCKS = [("C", "D", "E", "F"), ("H", "I", "J", "K"), ("M", "N", "O", "P")]
S2_M0, S2_MONTHS = 6, 18
S2_TOTAL = S2_M0 + S2_MONTHS          # 24
S2_CHK_IN, S2_CHK_OUT, S2_GAP, S2_START = S2_TOTAL + 2, S2_TOTAL + 3, S2_TOTAL + 4, S2_TOTAL + 5

ARIAL = "Arial"
FMT = {
    "usd": '$#,##0;[Red]-$#,##0;"–"',
    "usd2": '$#,##0.00;[Red]-$#,##0.00;"–"',
    "pct": "0.0%",
    "ppl": "#,##0",
    "dec": "#,##0.0",
    "x": '0.00"x";[Red]-0.00"x";"–"',
    "uzs": '#,##0" сум"',
    "int": "0",
    "text": "@",
    "calc_text": "General",
}
UNIT = {"usd": "$", "usd2": "$", "pct": "%", "ppl": "чел", "dec": "шт", "x": "x", "uzs": "сум", "int": "мес",
        "text": "", "calc_text": ""}


def font(size=10, **kw):
    return Font(name=ARIAL, size=size, **kw)


FILL = {k: PatternFill("solid", fgColor=v) for k, v in {
    "input": "FFF2CC", "header": "1F4E79", "total": "DDEBF7", "sub": "F2F2F2", "colhead": "D9D9D9",
    "summary": "E2EFDA", "summary_total": "C6E0B4", "scen": "D9E1F2"}.items()}
THIN = Side(style="thin", color="BFBFBF")

ROWS = []


def H(label):
    ROWS.append(dict(kind="header", label=label))


def SUB(label):
    ROWS.append(dict(kind="subheader", label=label))


def BL():
    ROWS.append(dict(kind="blank"))


def I(key, label, fmt, note="", unit=None, valid=None):  # noqa: E743
    ROWS.append(dict(kind="input", key=key, label=label, fmt=fmt, note=note, unit=unit, valid=valid or fmt))


def F(key, label, fmt, formula, note="", unit=None, style=None):
    ROWS.append(dict(kind="formula", key=key, label=label, fmt=fmt, formula=formula, note=note, unit=unit, style=style))


# ---------------------------------------------------------------- Лист 1: строки
H("ИТОГ ЦИКЛА")
F("s_rev", "Выручка по договорам: холодный + тёплый запуск", "usd", "={rev_total}",
  "Сколько продали в деньгах — до комиссий и рассрочек", style="summary")
F("s_losses", "Съели эквайринг, комиссия рассрочки, невыплаты, возвраты", "usd", "=-{losses_total}",
  "Расшифровка — в блоке 7", style="summary")
F("s_net", "Чистые деньги за цикл", "usd", "={net_money}", "Выручка минус потери", style="summary_total")
F("s_now", "   из них пришло сразу, в месяц продаж", "usd", "={cash_now_total}",
  "Полная оплата, партнёрская рассрочка, первый платёж своей рассрочки — минус возвраты", style="summary")
F("s_tail", "   из них придёт позже по своей рассрочке", "usd", "={cash_tail_total}",
  "Ежемесячные платежи учеников — по месяцам на листе «Деньги по месяцам»", style="summary")
F("s_costs", "Расходы цикла", "usd", "=-{costs_total}",
  "Реклама, разовые вложения, ИИ, менеджер, команда, сервисы, кураторы, налог", style="summary")
F("s_profit", "ПРИБЫЛЬ ЦИКЛА", "usd", "={profit_pre}", "Чистые деньги минус расходы", style="summary_total")
F("s_profit_rec", "Прибыль без разовых вложений", "usd", "={profit_recurring}",
  "Сколько даст повторный цикл, когда ИИ-продавец уже настроен", style="summary")
F("s_profit_school", "Прибыль школы после доли эксперта / партнёра", "usd", "={profit_school}", "", style="summary")
F("s_owner", "Собственнику после налога на дивиденды", "usd", "={owner_money}",
  "Если выводить прибыль из ООО нерезиденту", style="summary")
F("s_romi", "ROMI: прибыль ÷ реклама", "x", "={romi}", "Сколько прибыли на 1 $ рекламы", style="summary")
F("s_sales", "Продаж за цикл", "dec", "={sales_total}", "Холодный + тёплый запуск", style="summary")
F("s_cpl", "Цена холодной регистрации сейчас", "usd2", "={cpl}", "Из блока 1", style="summary")
F("s_cpl_max", "Предельная цена регистрации", "usd2", "={cpl_be_cycle}",
  "Дороже этой цены реклама не отбивается, даже без учёта команды и окладов", style="summary")
F("s_be_budget", "Рекламный бюджет, чтобы выйти в ноль", "usd", "={be_budget}",
  "При текущих конверсиях. Приближённо — см. блок 9", style="summary")
F("s_start_cash", "Нужно вложить до первых продаж", "usd", "={start_cash}",
  "Реклама + разовые вложения + команда + первый месяц окладов и сервисов", style="summary")
F("s_gap", "Самая низкая точка кассы по месяцам", "usd", "={S2_GAP}",
  "Минимум нарастающего итога на листе «Деньги по месяцам». Внутри месяца 1 разрыв больше — см. строку выше",
  style="summary")
BL()

H("0. ПРОДУКТ И ЦЕНЫ")
I("price_cold", "Цена курса в холодном запуске", "usd",
  "Чек 300 $ — по ответу заказчика. У конкурентов: фриланс 170–595 $, ИИ-академии 255–510 $; база брифа — 212 $")
I("price_warm", "Цена курса в тёплом запуске", "usd", "Можно поставить спеццену для подписчиков канала")
I("upsell_rate", "Апсейл: докупили старший тариф или сопровождение, % покупателей", "pct",
  "Рычаг «допродать». По умолчанию 0 — впиши, если будет старший тариф")
I("upsell_price", "Цена апсейла (доплата)", "usd")
I("fx", "Курс: сум за 1 $ (для справки)", "uzs", "ЦБ РУз на 15.09.2026 — 11 765 сум", unit="сум", valid="nonneg")
F("price_cold_uzs", "Цена курса в сумах", "uzs", "={price_cold}*{fx}", "Для объявлений и лендинга", unit="сум")
BL()

H("1. ХОЛОДНЫЙ ТРАФИК — таргет Instagram / Facebook")
I("budget", "Рекламный бюджет на холодный запуск", "usd",
  "Главный рычаг: поменяй — пересчитается весь цикл. Для Telegram-посевов подбери CPM, CTR и конверсию под свою цену регистрации",
  valid="nonneg")
I("cpm", "CPM — цена 1 000 показов", "usd2", "Meta в УЗ: 1–3 $ на широкую аудиторию, в Instagram 3–8 $", valid="nonneg")
F("impr", "Показы", "ppl", "=IF({cpm}>0,{budget}/{cpm}*1000,0)", "Бюджет ÷ CPM × 1 000", unit="шт")
I("ctr", "CTR — кликнули на рекламу", "pct", "Хорошим в Instagram УЗ считают 1,5–3%")
F("clicks", "Клики", "ppl", "={impr}*{ctr}", "Показы × CTR", unit="шт")
F("cpc", "Цена клика (CPC)", "usd2", "=IF({clicks}>0,{budget}/{clicks},0)", "Бюджет ÷ клики")
I("cr_reg", "Конверсия клик → регистрация (старт бота)", "pct", "Со страницы регистрации регистрируются 80–90%, с Telegram Ads на бота — 19–21% кликнувших; с Meta через лендинг и бота ниже")
F("reg", "Регистрации", "ppl", "={clicks}*{cr_reg}", "Клики × конверсия", style="sub")
F("cpl", "Цена регистрации (CPL)", "usd2", "=IF({reg}>0,{budget}/{reg},0)",
  "Бюджет ÷ регистрации. Лид в образовании УЗ — 2,3–3,6 $, в сезон июль–сентябрь — 10–12 $; база брифа — 6 $", style="sub")
BL()

H("2. ХОЛОДНЫЙ ЭФИР — вебинар или марафон")
I("cold_format", "Формат (метка для себя)", "text", "Выбери из списка. На расчёт не влияет — меняй проценты ниже",
  valid="list_format")
I("att1", "Дошли на эфир (день 1), % регистраций", "pct", "В среднем 15–20% (РФ); 20–30% — норма, 50% — хорошо (гайд СНГ); база брифа — 35%")
I("att2", "Досидели до продажи (день 2), % пришедших", "pct",
  "Для вебинара — 100%. Для марафона — доля пришедших, кто дошёл до продающего дня")
F("part", "Участники продающего эфира", "ppl", "={reg}*{att1}*{att2}", "Регистрации × доходимость × удержание", style="sub")
I("ap_rate", "Автооплаты на эфире, % участников", "pct", "Оплатили по ссылке во время эфира")
F("ap", "Автооплаты", "dec", "={part}*{ap_rate}")
I("bk_rate", "Брони / предоплаты, % участников", "pct", "Внесли часть денег, дальше их ведёт менеджер")
F("bk", "Брони", "dec", "={part}*{bk_rate}")
I("or_rate", "Заявки без оплаты, % участников", "pct", "Оставили контакты — дальше звонит менеджер")
F("orr", "Заявки", "dec", "={part}*{or_rate}")
F("chk_cold", "Проверка: автооплаты + брони + заявки не больше 100%", "calc_text",
  '=IF({ap_rate}+{bk_rate}+{or_rate}>1,"ОШИБКА: больше 100%","ок")', style="check")
F("noact", "Были на эфире, но ничего не сделали", "ppl", "=MAX(0,{part}-{ap}-{bk}-{orr})", "Их дожимает ИИ-продавец в боте")
F("noshow", "Не дошли до продажи", "ppl", "=MAX(0,{reg}-{part})", "Им бот шлёт запись, дальше пишет ИИ-продавец")
BL()

H("3. КАК ПРОДАЮТ ИИ-ПРОДАВЕЦ И МЕНЕДЖЕР — общие ставки для обоих запусков")
I("bot_share", "В боте, % регистраций", "pct", "Если регистрация идёт через Telegram-бота — 100%")
I("ai_eng_att", "ИИ: ответили в диалоге — из «были, но ничего не сделали», %", "pct",
  "Бот пишет после эфира; считаем тех, кто ответил")
I("ai_eng_no", "ИИ: ответили в диалоге — из «не дошли», %", "pct", "Холоднее: не видели продажу")
I("ai_pay", "ИИ: оплатили сами по ссылке, % диалогов", "pct", "ИИ закрывает без человека")
I("ai_hot", "ИИ: передал менеджеру горячего, % диалогов", "pct", "Готов купить, но нужен звонок или рассрочка")
F("chk_ai", "Проверка: оплатили + горячие не больше 100%", "calc_text",
  '=IF({ai_pay}+{ai_hot}>1,"ОШИБКА: больше 100%","ок")', style="check")
I("m_bk", "Менеджер: доводит бронь до оплаты, %", "pct", "Заказ → оплата на GetCourse — 66,8% в пик, в примере финмодели — 46%, в исходном файле — 80%")
I("m_or", "Менеджер: доводит заявку до оплаты, %", "pct", "Звонок → оплата 50–55% по гайду (не проверено); база брифа — 40%; в исходном файле — 40–50%")
I("m_hot", "Менеджер: доводит горячего от ИИ до оплаты, %", "pct", "Лид уже прогрет ИИ — оценка")
BL()

H("4. ХОЛОДНЫЙ ЗАПУСК — ПРОДАЖИ")
F("c_ap", "Автооплаты на эфире", "dec", "={ap}")
F("c_m_bk", "Менеджер: оплаты по броням", "dec", "={bk}*{m_bk}", "Брони × доводимость")
F("c_m_or", "Менеджер: оплаты по заявкам", "dec", "={orr}*{m_or}", "Заявки × доводимость")
F("c_dialogs", "ИИ: диалогов в боте", "ppl", "=({noact}*{ai_eng_att}+{noshow}*{ai_eng_no})*{bot_share}",
  "(ничего не сделали × % ответа + не дошли × % ответа) × % в боте", unit="шт")
F("c_ai_pay", "ИИ: оплатили сами", "dec", "={c_dialogs}*{ai_pay}")
F("c_hot", "ИИ → менеджер: горячих лидов", "dec", "={c_dialogs}*{ai_hot}")
F("c_m_hot", "Менеджер: оплаты от горячих лидов ИИ", "dec", "={c_hot}*{m_hot}")
F("c_sales", "ИТОГО продаж холодного запуска", "dec", "={c_ap}+{c_m_bk}+{c_m_or}+{c_ai_pay}+{c_m_hot}", style="total")
F("c_upsell", "Апсейлов", "dec", "={c_sales}*{upsell_rate}")
F("c_rev", "Выручка холодного запуска", "usd", "={c_sales}*{price_cold}+{c_upsell}*{upsell_price}",
  "Продажи × цена + апсейлы × доплата", style="total")
F("c_mgr_rev", "   в т.ч. через менеджера (с апсейлом)", "usd",
  "=({c_m_bk}+{c_m_or}+{c_m_hot})*{price_cold}+{c_upsell}*{upsell_price}", "С этой суммы менеджер получает процент")
F("c_ai_rev", "   в т.ч. с участием ИИ: сам + горячие", "usd", "=({c_ai_pay}+{c_m_hot})*{price_cold}")
F("c_conv", "Конверсия регистрация → покупка", "pct", "=IF({reg}>0,{c_sales}/{reg},0)")
F("c_cac", "Цена покупателя: реклама ÷ продажи", "usd", "=IF({c_sales}>0,{budget}/{c_sales},0)")
F("c_roas", "ROAS холодного: выручка ÷ реклама", "x", "=IF({budget}>0,{c_rev}/{budget},0)",
  "Это «окупаемость» из старого файла — без учёта расходов")
BL()

H("5. TELEGRAM-КАНАЛ — копим базу к тёплому запуску")
I("ch_start", "Подписчиков в канале на старте цикла", "ppl",
  "Если у эксперта уже есть канал или блог — впиши. Это главная ячейка для тёплого запуска", valid="nonneg")
I("ch_sub", "Подписались на канал из холодного запуска, % регистраций", "pct", "Бот зовёт в канал после регистрации и эфира")
F("ch_from_cold", "Пришли в канал из холодного запуска", "ppl", "={reg}*{ch_sub}")
I("ch_org", "Органический прирост до тёплого запуска", "ppl", "Без рекламы", valid="nonneg")
I("ch_churn", "Отписки до тёплого запуска, %", "pct", "Отписки 8–12% (РФ, период не указан)")
F("ch_churn_n", "Отписались", "ppl", "=({ch_start}+{ch_from_cold}+{ch_org})*{ch_churn}")
F("ch_buyers", "Уже купили курс — не целевые", "ppl", "={c_sales}*{ch_sub}", "Покупатели холодного запуска, которые есть в канале")
F("ch_base", "Целевая база канала к тёплому запуску", "ppl",
  "=MAX(0,{ch_start}+{ch_from_cold}+{ch_org}-{ch_churn_n}-{ch_buyers})", style="sub")
I("gap", "Через сколько месяцев после холодного — тёплый запуск", "int", "По плану — раз в 2 месяца. От 1 до 6",
  unit="мес", valid="gap")
BL()

H("6. ТЁПЛЫЙ ЗАПУСК — марафон 2 дня в канале")
I("warm_format", "Формат (метка для себя)", "text", "На расчёт не влияет", valid="list_format")
I("w_reg_ch_rate", "Зарегистрировались из канала, % целевой базы", "pct",
  "Анонс идёт серией постов — берём долю подписчиков, а не охват одного поста")
F("w_reg_ch", "Регистрации из канала", "ppl", "={ch_base}*{w_reg_ch_rate}")
F("w_bot_rest", "База бота вне канала, кто не купил", "ppl", "=MAX(0,({reg}*{bot_share}-{c_sales})*(1-{ch_sub}))",
  "Им шлём рассылку через бота")
I("w_reg_bot_rate", "Зарегистрировались из рассылки бота, %", "pct")
F("w_reg_bot", "Регистрации из бота", "ppl", "={w_bot_rest}*{w_reg_bot_rate}")
I("w_budget", "Бюджет на посевы / рекламу тёплого запуска", "usd", "По умолчанию 0 — тёплый запуск на своей базе", valid="nonneg")
I("w_cpl", "Цена регистрации с посевов", "usd2", "Подписчик в Telegram Ads в УЗ — 0,15–1,5 €", valid="nonneg")
F("w_reg_paid", "Регистрации с посевов", "ppl", "=IF({w_cpl}>0,{w_budget}/{w_cpl},0)")
F("w_reg", "Регистраций на тёплый запуск всего", "ppl", "={w_reg_ch}+{w_reg_bot}+{w_reg_paid}", style="sub")
I("w_att1", "Дошли в день 1, % регистраций", "pct", "Холодный марафон — 25%, тёплый — до 68% (РФ)")
I("w_att2", "Дошли до продающего дня 2, % пришедших", "pct", "В тёплом марафоне на день 2 вернулись 72% (РФ)")
F("w_part", "Участники продающего эфира", "ppl", "={w_reg}*{w_att1}*{w_att2}", style="sub")
I("w_ap_rate", "Автооплаты, % участников", "pct")
F("w_ap", "Автооплаты", "dec", "={w_part}*{w_ap_rate}")
I("w_bk_rate", "Брони / предоплаты, % участников", "pct")
F("w_bk", "Брони", "dec", "={w_part}*{w_bk_rate}")
I("w_or_rate", "Заявки без оплаты, % участников", "pct")
F("w_or", "Заявки", "dec", "={w_part}*{w_or_rate}")
F("chk_warm", "Проверка: автооплаты + брони + заявки не больше 100%", "calc_text",
  '=IF({w_ap_rate}+{w_bk_rate}+{w_or_rate}>1,"ОШИБКА: больше 100%","ок")', style="check")
F("w_noact", "Были на эфире, но ничего не сделали", "ppl", "=MAX(0,{w_part}-{w_ap}-{w_bk}-{w_or})")
F("w_noshow", "Не дошли до продажи", "ppl", "=MAX(0,{w_reg}-{w_part})")
F("w_m_bk", "Менеджер: оплаты по броням", "dec", "={w_bk}*{m_bk}", "Ставки менеджера и ИИ — из блока 3")
F("w_m_or", "Менеджер: оплаты по заявкам", "dec", "={w_or}*{m_or}")
F("w_dialogs", "ИИ: диалогов в боте", "ppl", "=({w_noact}*{ai_eng_att}+{w_noshow}*{ai_eng_no})*{bot_share}", unit="шт")
F("w_ai_pay", "ИИ: оплатили сами", "dec", "={w_dialogs}*{ai_pay}")
F("w_hot", "ИИ → менеджер: горячих лидов", "dec", "={w_dialogs}*{ai_hot}")
F("w_m_hot", "Менеджер: оплаты от горячих лидов ИИ", "dec", "={w_hot}*{m_hot}")
F("w_sales", "ИТОГО продаж тёплого запуска", "dec", "={w_ap}+{w_m_bk}+{w_m_or}+{w_ai_pay}+{w_m_hot}", style="total")
F("w_upsell", "Апсейлов", "dec", "={w_sales}*{upsell_rate}")
F("w_rev", "Выручка тёплого запуска", "usd", "={w_sales}*{price_warm}+{w_upsell}*{upsell_price}", style="total")
F("w_mgr_rev", "   в т.ч. через менеджера (с апсейлом)", "usd",
  "=({w_m_bk}+{w_m_or}+{w_m_hot})*{price_warm}+{w_upsell}*{upsell_price}")
F("w_ai_rev", "   в т.ч. с участием ИИ: сам + горячие", "usd", "=({w_ai_pay}+{w_m_hot})*{price_warm}")
F("w_conv", "Конверсия регистрация → покупка", "pct", "=IF({w_reg}>0,{w_sales}/{w_reg},0)")
BL()

H("7. ОПЛАТЫ, РАССРОЧКИ И КОМИССИИ — на все продажи цикла")
I("pay_full", "Полная оплата, % продаж", "pct", "Картой через Payme / Click / Uzum")
I("pay_partner", "Партнёрская рассрочка (Uzum Nasiya, Alif), % продаж", "pct",
  "Сервис платит школе сразу и забирает комиссию. Рассрочка 0% на 6–12 мес — стандарт рынка; Alif одобряет 50–60% заявок, Uzum — до 97%")
F("pay_own", "Своя рассрочка школы, % продаж", "pct", "=MAX(0,1-{pay_full}-{pay_partner})", "Остаток до 100%")
F("chk_pay", "Проверка: полная + партнёрская не больше 100%", "calc_text",
  '=IF({pay_full}+{pay_partner}>1,"ОШИБКА: больше 100%","ок")', style="check")
I("acq", "Эквайринг (Payme, Click, Uzum), % с платежа", "pct", "Payme ~1–1,5%, Click ~1–2%, Uzum 0,4–1,5%")
I("partner_fee", "Комиссия партнёрской рассрочки для школы, %", "pct",
  "Ставки для школ не найдены — запросить у Uzum, Alif, Anorbank. Если 0% для покупателя, наценку платит школа. Ориентир РФ: 8,5% за 3 мес, 15% за 12 мес")
I("own_term", "Своя рассрочка: срок, месяцев", "int", "Первый платёж — в месяц покупки, остальные — каждый следующий месяц. От 1 до 12",
  unit="мес", valid="term")
F("own_n", "Срок в расчёте", "int", "=MAX(1,MIN(12,ROUND({own_term},0)))", "Защита от 0 и дробей", unit="мес")
I("own_first", "Своя рассрочка: первый платёж, % цены", "pct", "Остаток делится поровну на следующие месяцы")
F("own_first_eff", "Первый платёж в расчёте", "pct", "=IF({own_n}=1,1,{own_first})", "При сроке 1 месяц — вся сумма сразу")
I("own_default", "Своя рассрочка: не доплатят, % от остатка", "pct", "Данных нет; оценка ресёрча — 15%")
I("refund", "Возвраты, % выручки", "pct", "Допущение брифа — 5%")
SUB("Холодный запуск")
F("c_now_gross", "Пришло в месяц продаж", "usd",
  "={c_rev}*({pay_full}*(1-{acq})+{pay_partner}*(1-{partner_fee})+{pay_own}*{own_first_eff}*(1-{acq}))",
  "Полная оплата − эквайринг + партнёрская − комиссия + первый платёж своей рассрочки")
F("c_tail", "Придёт позже по своей рассрочке", "usd",
  "={c_rev}*{pay_own}*(1-{own_first_eff})*(1-{own_default})*(1-{acq})", "Остаток минус невыплаты и эквайринг")
F("c_own_monthly", "   это по … $ в месяц", "usd", "=IF({own_n}>1,{c_tail}/({own_n}-1),0)", "Со 2-го месяца, всего «срок − 1» платежей")
F("c_refund", "Возвраты", "usd", "={c_rev}*{refund}")
F("c_net", "Чистые деньги холодного запуска", "usd", "={c_now_gross}+{c_tail}-{c_refund}", style="sub")
SUB("Тёплый запуск")
F("w_now_gross", "Пришло в месяц продаж", "usd",
  "={w_rev}*({pay_full}*(1-{acq})+{pay_partner}*(1-{partner_fee})+{pay_own}*{own_first_eff}*(1-{acq}))")
F("w_tail", "Придёт позже по своей рассрочке", "usd",
  "={w_rev}*{pay_own}*(1-{own_first_eff})*(1-{own_default})*(1-{acq})")
F("w_own_monthly", "   это по … $ в месяц", "usd", "=IF({own_n}>1,{w_tail}/({own_n}-1),0)")
F("w_refund", "Возвраты", "usd", "={w_rev}*{refund}")
F("w_net", "Чистые деньги тёплого запуска", "usd", "={w_now_gross}+{w_tail}-{w_refund}", style="sub")
SUB("Весь цикл")
F("rev_total", "Выручка по договорам", "usd", "={c_rev}+{w_rev}", style="total")
F("loss_acq", "Потери: эквайринг", "usd",
  "={rev_total}*({pay_full}+{pay_own}*{own_first_eff}+{pay_own}*(1-{own_first_eff})*(1-{own_default}))*{acq}")
F("loss_fee", "Потери: комиссия партнёрской рассрочки", "usd", "={rev_total}*{pay_partner}*{partner_fee}")
F("loss_default", "Потери: не доплатили по своей рассрочке", "usd", "={rev_total}*{pay_own}*(1-{own_first_eff})*{own_default}")
F("loss_refund", "Потери: возвраты", "usd", "={c_refund}+{w_refund}")
F("losses_total", "Потери всего", "usd", "={loss_acq}+{loss_fee}+{loss_default}+{loss_refund}", style="sub")
F("loss_share", "Потери, % выручки", "pct", "=IF({rev_total}>0,{losses_total}/{rev_total},0)")
F("net_money", "Чистые деньги за цикл", "usd", "={c_net}+{w_net}", style="total")
F("cash_now_total", "   пришло сразу, после возвратов", "usd", "={c_now_gross}-{c_refund}+{w_now_gross}-{w_refund}")
F("cash_tail_total", "   придёт позже по своей рассрочке", "usd", "={c_tail}+{w_tail}")
F("chk_money", "Сверка: выручка − потери = чистые деньги", "calc_text",
  '=IF(ABS({rev_total}-{losses_total}-{net_money})<0.01,"ок","ОШИБКА")', style="check")
BL()

H("8. РАСХОДЫ ЦИКЛА")
I("ai_setup", "ИИ-продавец: разработка и настройка, разово", "usd",
  "В УЗ — 5–25 млн сум; 1 300 $ ≈ 15 млн сум. Если собираете сами — поставь 0", valid="nonneg")
I("one_time_other", "Прочие разовые вложения: ООО, юрист, запись курса, лендинг, бот", "usd",
  "По умолчанию 0. Пошлина за ООО ≈ 25 $; принимать оплату через Payme/Click можно только своему ООО", valid="nonneg")
I("ai_sub", "ИИ-продавец: ИИ и хостинг, $ в месяц", "usd", "В УЗ — 0,3–3 млн сум в месяц", valid="nonneg")
I("ai_dialog_cost", "ИИ-продавец: доплата за 1 диалог", "usd2", "Сервисы берут 5–30 ₽ за диалог; свой бот дешевле", valid="nonneg")
I("m_salary", "Менеджер: оклад, $ в месяц", "usd", "Оценка. Можно работать на проценте без оклада", valid="nonneg")
I("m_pct", "Менеджер: % с выручки своих сделок", "pct", "В брифе затраты на продажи — 10% выручки")
I("team", "Команда запуска на цикл: таргетолог, техспец, дизайнер, монтаж", "usd", "Оценка — уточнить у команды", valid="nonneg")
I("services", "Сервисы: бот-платформа, стрим, лендинг, LMS, $ в месяц", "usd", "Оценка", valid="nonneg")
I("curator", "Обучение: кураторы и поддержка на 1 ученика", "usd", "Допущение брифа: 10–25 $ в зависимости от чека", valid="nonneg")
I("tax", "Налог, % от поступлений", "pct", "ООО — налог с оборота 4% до 1 млрд сум в год; ИП и самозанятые — 1%. Проверить у бухгалтера")
I("partner_share", "Доля эксперта / партнёра, % от прибыли", "pct", "Если делите прибыль с экспертом. По умолчанию 0")
I("div_tax", "Налог на дивиденды при выводе прибыли, %", "pct", "Нерезиденту — 10%, для России по договору тоже 10%")
F("months", "Месяцев в цикле", "int", "={gap}+1", "Тёплый запуск идёт в последний месяц цикла", unit="мес")
F("sales_total", "Учеников за цикл", "dec", "={c_sales}+{w_sales}")
F("cost_ads", "Реклама: холодная + посевы тёплого", "usd", "={budget}+{w_budget}")
F("cost_one_time", "Разовые вложения", "usd", "={ai_setup}+{one_time_other}")
F("cost_ai", "ИИ-продавец: ИИ, хостинг и диалоги", "usd", "={ai_sub}*{months}+({c_dialogs}+{w_dialogs})*{ai_dialog_cost}")
F("cost_mgr", "Менеджер: оклад + процент", "usd", "={m_salary}*{months}+({c_mgr_rev}+{w_mgr_rev})*{m_pct}")
F("cost_team", "Команда запуска", "usd", "={team}")
F("cost_services", "Сервисы", "usd", "={services}*{months}")
F("cost_curators", "Кураторы", "usd", "={sales_total}*{curator}")
F("cost_tax", "Налог", "usd", "={net_money}*{tax}",
  "С чистых денег за цикл, включая будущие платежи рассрочки. Упрощение: налоговая база до комиссий сервисов чуть больше")
F("costs_total", "РАСХОДЫ ВСЕГО", "usd",
  "={cost_ads}+{cost_one_time}+{cost_ai}+{cost_mgr}+{cost_team}+{cost_services}+{cost_curators}+{cost_tax}", style="total")
SUB("Для кассы по месяцам")
F("c_var_costs", "Расходы в месяц холодного запуска", "usd",
  "={budget}+{c_dialogs}*{ai_dialog_cost}+{c_mgr_rev}*{m_pct}+{c_sales}*{curator}",
  "Реклама, диалоги ИИ, % менеджера, кураторы (разовые вложения — тоже в месяц 1)")
F("w_var_costs", "Расходы в месяц тёплого запуска", "usd",
  "={w_budget}+{w_dialogs}*{ai_dialog_cost}+{w_mgr_rev}*{m_pct}+{w_sales}*{curator}")
F("fixed_monthly", "Постоянные расходы в месяц", "usd", "={ai_sub}+{m_salary}+{services}+{team}/{months}",
  "Команда разложена поровну на месяцы цикла")
F("start_cash", "Нужно вложить до первых продаж", "usd",
  "={budget}+{cost_one_time}+{team}+{ai_sub}+{m_salary}+{services}", "Реклама, разовые вложения, команда и первый месяц окладов и сервисов", style="sub")
BL()

H("9. ИТОГ И ЮНИТ-ЭКОНОМИКА")
F("profit_pre", "Прибыль цикла", "usd", "={net_money}-{costs_total}", style="total")
F("profit_recurring", "Прибыль без разовых вложений", "usd", "={profit_pre}+{cost_one_time}",
  "Повторный цикл без затрат на настройку ИИ")
F("partner_cut", "Доля эксперта / партнёра", "usd", "=MAX(0,{profit_pre})*{partner_share}")
F("profit_school", "Прибыль школы", "usd", "={profit_pre}-{partner_cut}", style="sub")
F("owner_money", "Собственнику после налога на дивиденды", "usd", "={profit_school}-MAX(0,{profit_school})*{div_tax}",
  "Налог берём только с положительной прибыли")
F("margin", "Маржинальность: прибыль ÷ выручка", "pct", "=IF({rev_total}>0,{profit_pre}/{rev_total},0)")
F("romi", "ROMI: прибыль ÷ реклама", "x", "=IF({cost_ads}>0,{profit_pre}/{cost_ads},0)")
F("warm_share", "Доля тёплого запуска в выручке", "pct", "=IF({rev_total}>0,{w_rev}/{rev_total},0)",
  "В первом цикле с новым каналом маленькая — база канала ещё не накоплена")
F("ai_rev_total", "Выручка с участием ИИ-продавца", "usd", "={c_ai_rev}+{w_ai_rev}")
F("ai_roi", "Окупаемость ИИ: выручка с ИИ ÷ расходы на ИИ", "x",
  "=IF({cost_ai}+{ai_setup}>0,{ai_rev_total}/({cost_ai}+{ai_setup}),0)", "Считаем вместе с разовой настройкой")
F("mgr_rev_total", "Выручка через менеджера", "usd", "={c_mgr_rev}+{w_mgr_rev}")
F("mgr_share", "Доля менеджера в выручке", "pct", "=IF({rev_total}>0,{mgr_rev_total}/{rev_total},0)")
SUB("На 1 холодную регистрацию")
F("u_rev_cold", "Выручка холодного запуска на 1 регистрацию", "usd2", "=IF({reg}>0,{c_rev}/{reg},0)")
F("u_rev_cycle", "Выручка цикла на 1 регистрацию, с тёплым", "usd2", "=IF({reg}>0,{rev_total}/{reg},0)")
F("money_ratio", "Доходит деньгами, % выручки", "pct", "=IF({rev_total}>0,{net_money}/{rev_total},0)")
F("u_contrib_cold", "Остаётся с 1 регистрации — только холодный", "usd2",
  "=IF({reg}>0,({c_rev}*{money_ratio}*(1-{tax})-{c_dialogs}*{ai_dialog_cost}-{c_mgr_rev}*{m_pct}-{c_sales}*{curator})/{reg},0)",
  "После комиссий, налога, диалогов ИИ, % менеджера и кураторов. До рекламы и постоянных расходов")
F("cpl_be_cycle", "Предельная цена регистрации — с тёплым запуском", "usd2",
  "=IF({reg}>0,({rev_total}*{money_ratio}*(1-{tax})-({c_dialogs}+{w_dialogs})*{ai_dialog_cost}-({c_mgr_rev}+{w_mgr_rev})*{m_pct}-{sales_total}*{curator})/{reg},0)",
  "Дороже — реклама не отбивается даже без учёта постоянных расходов", style="sub")
F("cpl_headroom", "Запас: трафик может подорожать в … раз", "x", "=IF({cpl}>0,{cpl_be_cycle}/{cpl},0)")
F("fixed_costs", "Постоянные и разовые расходы цикла", "usd",
  "={team}+({ai_sub}+{m_salary}+{services})*{months}+{w_budget}+{cost_one_time}",
  "Команда, оклад, ИИ, сервисы, посевы тёплого, разовые вложения")
F("be_reg", "Сколько холодных регистраций нужно, чтобы выйти в ноль", "ppl",
  '=IF({cpl_be_cycle}-{cpl}>0,{fixed_costs}/({cpl_be_cycle}-{cpl}),"не выходим")',
  "Приближённо: считаем, что всё, кроме стартовой базы канала и органики, растёт вместе с трафиком")
F("be_budget", "Рекламный бюджет, чтобы выйти в ноль", "usd", '=IF(ISNUMBER({be_reg}),{be_reg}*{cpl},"не выходим")',
  "Регистрации для нуля × текущая цена регистрации")
F("chk_unit", "Сверка: регистрации × (предельная − текущая цена) − постоянные = прибыль", "calc_text",
  '=IF({reg}=0,"—",IF(ABS({reg}*({cpl_be_cycle}-{cpl})-{fixed_costs}-{profit_pre})<0.5,"ок","ОШИБКА"))', style="check")


# ---------------------------------------------------------------- сборка книги
TOKEN = re.compile(r"\{(\w+)\}")


def layout():
    rowmap, r = {}, FIRST_ROW
    for row in ROWS:
        row["row"] = r
        if row.get("key"):
            if row["key"] in rowmap:
                raise ValueError(f"duplicate key {row['key']}")
            rowmap[row["key"]] = r
        r += 1
    return rowmap


def resolve(formula, sc, rowmap):
    def rep(m):
        k = m.group(1)
        if k == "S2_GAP":
            return f"'{S2}'!{S2_BLOCKS[sc][3]}{S2_GAP}"
        if k not in rowmap:
            raise KeyError(f"unknown token {k} in {formula}")
        return f"{COLS[sc]}{rowmap[k]}"
    return TOKEN.sub(rep, formula)


def s1ref(key, sc, rowmap):
    return f"'{S1}'!${COLS[sc]}${rowmap[key]}"


def style_row(ws, r, fill=None, bold=False, color=None, italic=False, size=10):
    for col in "BCDEFGH":
        c = ws[f"{col}{r}"]
        if fill:
            c.fill = fill
        c.font = font(size=size, bold=bold, italic=italic, color=color)


def build_sheet1(wb, rowmap):
    ws = wb.active
    ws.title = S1
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 62, "C": 6, "D": 15, "E": 15, "F": 15, "G": 100, "H": 22}.items():
        ws.column_dimensions[col].width = w
    ws["B1"] = "Финмодель онлайн-школы в Узбекистане — один цикл"
    ws["B1"].font = font(size=15, bold=True)
    ws["B2"] = ("Холодный запуск с таргета → дожим ИИ-продавцом и менеджером → тёплый запуск в Telegram-канале. "
                "Все суммы — в долларах США.")
    ws["B2"].font = font(color="404040")
    ws["B3"] = "Жёлтые ячейки с синими цифрами — вводные: меняй их"
    ws["B3"].font = font(bold=True, color="0000FF")
    ws["B3"].fill = FILL["input"]
    ws["B4"] = ("Белые — формулы, не трогай. Три колонки — три сценария. Касса по месяцам — лист «Деньги по месяцам», "
                "откуда цифры — лист «Источники цифр», что было не так в старом файле — «Разбор исходника».")
    ws["B4"].font = font(color="404040")
    ws["B5"] = ("Контекст — бриф рынка от 15.09.2026 (research/01-brif-rynka). Решение «идти / не идти» не принято: "
                "модель показывает, при каких цифрах цикл окупается. Ссылка на источник каждой цифры — в колонке H.")
    ws["B5"].hyperlink = BRIEF_URL
    ws["B5"].font = font(color="0563C1", underline="single")
    headers = {"B": "Показатель", "C": "Ед.", "D": SCEN[0], "E": SCEN[1], "F": SCEN[2], "G": "Как считается / откуда цифра",
               "H": "Источник"}
    for col, text in headers.items():
        c = ws[f"{col}6"]
        c.value = text
        c.font = font(bold=True)
        c.fill = FILL["scen"] if col in COLS else FILL["colhead"]
        c.alignment = Alignment(horizontal="center" if col in "CDEF" else "left", vertical="center")
    ws.freeze_panes = "D7"
    src_urls = {key: urls for key, _label, _bench, _geo, urls, _date, _conf in SOURCES}

    dv = {
        "pct": DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False,
                              showErrorMessage=True, errorTitle="Процент", error="Введи значение от 0% до 100%"),
        "nonneg": DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True,
                                 errorTitle="Число", error="Введи число не меньше 0"),
        "gap": DataValidation(type="whole", operator="between", formula1="1", formula2="6", showErrorMessage=True,
                              errorTitle="Месяцы", error="Целое число от 1 до 6"),
        "term": DataValidation(type="whole", operator="between", formula1="1", formula2="12", showErrorMessage=True,
                               errorTitle="Месяцы", error="Целое число от 1 до 12"),
        "list_format": DataValidation(type="list", formula1='"Вебинар,Марафон 2 дня"', showErrorMessage=True,
                                      errorTitle="Формат", error="Выбери из списка"),
    }
    for v in dv.values():
        ws.add_data_validation(v)
    valid_map = {"usd": "nonneg", "usd2": "nonneg", "ppl": "nonneg", "pct": "pct"}

    red_fill = PatternFill("solid", bgColor="FFC7CE")
    for row in ROWS:
        r = row["row"]
        kind = row["kind"]
        if kind == "blank":
            continue
        if kind == "header":
            ws[f"B{r}"] = row["label"]
            style_row(ws, r, fill=FILL["header"], bold=True, color="FFFFFF", size=11)
            ws.row_dimensions[r].height = 20
            continue
        if kind == "subheader":
            ws[f"B{r}"] = row["label"]
            style_row(ws, r, bold=True, italic=True, color="1F4E79")
            continue
        ws[f"B{r}"] = row["label"]
        ws[f"C{r}"] = row["unit"] if row["unit"] is not None else UNIT[row["fmt"]]
        ws[f"G{r}"] = row["note"]
        src_key = row["key"] if row["key"] in src_urls else SOURCE_ALIAS.get(row["key"])
        h = ws[f"H{r}"]
        if src_key is not None and src_urls[src_key]:
            h.value = source_label(src_urls[src_key][0])
            h.hyperlink = src_urls[src_key][0]
            h.font = font(size=9, color="0563C1", underline="single")
        elif kind == "input" or src_key is not None:
            h.value = "оценка"
            h.font = font(size=9, color="595959", italic=True)
        style = row.get("style")
        base_fill = {"total": FILL["total"], "sub": FILL["sub"], "summary": FILL["summary"],
                     "summary_total": FILL["summary_total"]}.get(style)
        bold = style in ("total", "sub", "summary_total")
        for col in "BCG":
            c = ws[f"{col}{r}"]
            c.font = font(bold=bold and col == "B", color="595959" if col in "CG" else None,
                          italic=style == "check", size=9 if col == "G" else 10)
            if base_fill:
                c.fill = base_fill
        if base_fill:
            ws[f"H{r}"].fill = base_fill
        ws[f"C{r}"].alignment = Alignment(horizontal="center")
        ws[f"G{r}"].alignment = Alignment(indent=1)
        for sc, col in enumerate(COLS):
            c = ws[f"{col}{r}"]
            c.number_format = FMT[row["fmt"]]
            c.alignment = Alignment(horizontal="right")
            if kind == "input":
                c.value = DEF[row["key"]][sc]
                c.font = font(color="0000FF")
                c.fill = FILL["input"]
            else:
                c.value = resolve(row["formula"], sc, rowmap)
                c.font = font(bold=bold, italic=style == "check")
                if base_fill:
                    c.fill = base_fill
                if style == "check":
                    c.alignment = Alignment(horizontal="center")
        if kind == "input":
            vkey = row["valid"] if row["valid"] in dv else valid_map.get(row["valid"])
            if vkey:
                dv[vkey].add(f"D{r}:F{r}")
        if style == "check":
            ws.conditional_formatting.add(
                f"D{r}:F{r}", FormulaRule(formula=[f'ISNUMBER(SEARCH("ОШИБКА",D{r}))'], fill=red_fill,
                                           font=Font(name=ARIAL, color="9C0006", bold=True)))
        if style in ("summary_total",):
            ws.row_dimensions[r].height = 18
    return ws


def build_sheet2(wb, rowmap):
    ws = wb.create_sheet(S2)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 48
    for block in S2_BLOCKS:
        for col in block:
            ws.column_dimensions[col].width = 14
    ws.column_dimensions["G"].width = 2
    ws.column_dimensions["L"].width = 2
    ws["B1"] = "Деньги по месяцам — когда приходят и уходят деньги"
    ws["B1"].font = font(size=14, bold=True)
    ws["B2"] = ("Считается само из листа «Декомпозиция цикла», здесь ничего не вводится. Холодный запуск — месяц 1, "
                "тёплый — месяц «разрыв + 1». Своя рассрочка приходит частями в следующие месяцы.")
    ws["B2"].font = font(color="404040")
    heads = ["Поступления", "Расходы", "Итог месяца", "Нарастающим итогом"]
    for sc, block in enumerate(S2_BLOCKS):
        ws.merge_cells(f"{block[0]}4:{block[3]}4")
        c = ws[f"{block[0]}4"]
        c.value = SCEN[sc]
        c.font = font(bold=True, size=11)
        c.fill = FILL["scen"]
        c.alignment = Alignment(horizontal="center")
        for col, h in zip(block, heads):
            hc = ws[f"{col}5"]
            hc.value = h
            hc.font = font(bold=True)
            hc.fill = FILL["colhead"]
            hc.alignment = Alignment(horizontal="center", wrap_text=True)
    ws["B5"] = "Месяц"
    ws["B5"].font = font(bold=True)
    ws["B5"].fill = FILL["colhead"]
    ws.row_dimensions[5].height = 30
    ws.freeze_panes = "C6"
    money = FMT["usd"]

    for i in range(S2_MONTHS):
        r = S2_M0 + i
        ws[f"B{r}"] = i + 1
        ws[f"B{r}"].font = font()
        ws[f"B{r}"].alignment = Alignment(horizontal="left")
        for sc, (ci, co, cn, cc) in enumerate(S2_BLOCKS):
            def ref(key, _sc=sc):
                return s1ref(key, _sc, rowmap)
            m = f"$B{r}"
            ws[f"{ci}{r}"] = (
                f"=IF({m}=1,{ref('c_now_gross')}-{ref('c_refund')},0)"
                f"+IF(AND({m}>=2,{m}<={ref('own_n')}),{ref('c_own_monthly')},0)"
                f"+IF({m}={ref('months')},{ref('w_now_gross')}-{ref('w_refund')},0)"
                f"+IF(AND({m}>={ref('months')}+1,{m}<={ref('months')}+{ref('own_n')}-1),{ref('w_own_monthly')},0)"
            )
            ws[f"{co}{r}"] = (
                f"=IF({m}=1,{ref('c_var_costs')}+{ref('cost_one_time')},0)"
                f"+IF({m}={ref('months')},{ref('w_var_costs')},0)"
                f"+IF({m}<={ref('months')},{ref('fixed_monthly')},0)"
                f"+{ref('tax')}*{ci}{r}"
            )
            ws[f"{cn}{r}"] = f"={ci}{r}-{co}{r}"
            ws[f"{cc}{r}"] = f"={cn}{r}" if i == 0 else f"={cc}{r - 1}+{cn}{r}"
            for col in (ci, co, cn, cc):
                ws[f"{col}{r}"].number_format = money
                ws[f"{col}{r}"].font = font(bold=col == cc)
    last = S2_M0 + S2_MONTHS - 1
    ws[f"B{S2_TOTAL}"] = f"Итого за {S2_MONTHS} месяцев"
    ws[f"B{S2_CHK_IN}"] = "Сверка поступлений с листом 1 (должно быть 0)"
    ws[f"B{S2_CHK_OUT}"] = "Сверка расходов с листом 1 (должно быть 0)"
    ws[f"B{S2_GAP}"] = "Самая низкая точка кассы (минимум нарастающего итога)"
    ws[f"B{S2_START}"] = "Нужно вложить до первых продаж — внутри месяца 1"
    for r in (S2_TOTAL, S2_GAP, S2_START):
        ws[f"B{r}"].font = font(bold=True)
    for r in (S2_CHK_IN, S2_CHK_OUT):
        ws[f"B{r}"].font = font(italic=True, color="595959")
    for sc, (ci, co, cn, cc) in enumerate(S2_BLOCKS):
        for col in (ci, co, cn):
            ws[f"{col}{S2_TOTAL}"] = f"=SUM({col}{S2_M0}:{col}{last})"
            ws[f"{col}{S2_TOTAL}"].number_format = money
            ws[f"{col}{S2_TOTAL}"].font = font(bold=True)
            ws[f"{col}{S2_TOTAL}"].fill = FILL["total"]
        ws[f"{cc}{S2_TOTAL}"].fill = FILL["total"]
        ws[f"{ci}{S2_CHK_IN}"] = f"=ROUND({ci}{S2_TOTAL}-{s1ref('net_money', sc, rowmap)},2)"
        ws[f"{co}{S2_CHK_OUT}"] = f"=ROUND({co}{S2_TOTAL}-{s1ref('costs_total', sc, rowmap)},2)"
        ws[f"{cc}{S2_GAP}"] = f"=MIN(0,MIN({cc}{S2_M0}:{cc}{last}))"
        ws[f"{cc}{S2_START}"] = f"=-{s1ref('start_cash', sc, rowmap)}"
        for r, col in ((S2_CHK_IN, ci), (S2_CHK_OUT, co), (S2_GAP, cc), (S2_START, cc)):
            ws[f"{col}{r}"].number_format = FMT["usd2"] if r in (S2_CHK_IN, S2_CHK_OUT) else money
            ws[f"{col}{r}"].font = font(bold=r in (S2_GAP, S2_START), italic=r in (S2_CHK_IN, S2_CHK_OUT))
    return ws


def build_sheet3(wb):
    ws = wb.create_sheet(S3)
    ws.sheet_view.showGridLines = False
    for col, w in {"A": 2, "B": 5, "C": 30, "D": 22, "E": 44, "F": 58, "G": 50}.items():
        ws.column_dimensions[col].width = w
    ws["B1"] = "Разбор исходного файла «декомпозиция Узбекситан .xlsx»"
    ws["B1"].font = font(size=14, bold=True)
    ws["B2"] = ("Четыре листа — отдельные шаблоны из рублёвого запуска. Друг с другом они не связаны, трафик не связан "
                "с регистрациями, расходов и рассрочек нет, есть ошибка в формуле.")
    ws["B2"].font = font(color="404040")
    wrap = Alignment(wrap_text=True, vertical="top")
    r = 4
    ws[f"B{r}"] = "Что считает каждый лист"
    ws[f"B{r}"].font = font(bold=True, size=11, color="1F4E79")
    r += 1
    for col, text in zip("CDF", ["Лист", "Суть модели", "Главные цифры (как в файле)"]):
        ws[f"{col}{r}"] = text
    ws.merge_cells(f"D{r}:E{r}")
    ws.merge_cells(f"F{r}:G{r}")
    for col in "BCDEFG":
        ws[f"{col}{r}"].font = font(bold=True)
        ws[f"{col}{r}"].fill = FILL["colhead"]
    for sheet, essence, numbers in SHEETS_OVERVIEW:
        r += 1
        ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"] = sheet, essence, numbers
        ws.merge_cells(f"D{r}:E{r}")
        ws.merge_cells(f"F{r}:G{r}")
        for col in "CDF":
            ws[f"{col}{r}"].alignment = wrap
            ws[f"{col}{r}"].font = font(bold=col == "C")
        ws.row_dimensions[r].height = 48
    r += 2
    ws[f"B{r}"] = "Что не так и как исправлено в новой модели"
    ws[f"B{r}"].font = font(bold=True, size=11, color="1F4E79")
    r += 1
    for col, text in zip("BCDEFG", ["№", "Лист", "Ячейки", "Что там сейчас", "Почему это проблема", "Как в новой модели"]):
        ws[f"{col}{r}"] = text
        ws[f"{col}{r}"].font = font(bold=True)
        ws[f"{col}{r}"].fill = FILL["colhead"]
    for n, (sheet, cells, now, why, fix) in enumerate(FINDINGS, start=1):
        r += 1
        for col, text in zip("BCDEFG", [n, sheet, cells, now, why, fix]):
            c = ws[f"{col}{r}"]
            c.value = text
            c.alignment = wrap
            c.font = font(bold=col == "B")
        ws[f"B{r}"].alignment = Alignment(horizontal="center", vertical="top")
        ws.row_dimensions[r].height = 62
    ws.freeze_panes = "B4"
    return ws


def domain(url):
    netloc = urlparse(url).netloc
    return netloc[4:] if netloc.startswith("www.") else netloc


def source_label(url):
    return "бриф рынка" if url == BRIEF_URL else domain(url)


def build_sheet4(wb, rowmap):
    ws = wb.create_sheet(S4)
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 46, "C": 6, "D": 13, "E": 13, "F": 13, "G": 62, "H": 8,
              "I": 20, "J": 20, "K": 20, "L": 20, "M": 11, "N": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws["B1"] = "Источники цифр — откуда стартовые значения"
    ws["B1"].font = font(size=14, bold=True)
    ws["B2"] = ("Колонки «Минимум / Оптимум / Максимум» подтягиваются из модели. Бенчмарки собраны 15.09.2026; "
                "данных по конверсиям воронки в Узбекистане нет, поэтому часть взята по РФ — это отмечено в «Гео». "
                "Полное сырьё — research/04-finmodel-zapuska/бенчмарки-УЗ-2026-09.md.")
    ws["B2"].font = font(color="404040")
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("B2:I2")
    ws.row_dimensions[2].height = 32
    heads = ["Показатель", "Ед.", SCEN[0], SCEN[1], SCEN[2], "Что известно / бенчмарк", "Гео",
             "Источник 1", "Источник 2", "Источник 3", "Источник 4", "Дата", "Уверенность"]
    for col, text in zip("BCDEFGHIJKLMN", heads):
        c = ws[f"{col}4"]
        c.value = text
        c.font = font(bold=True)
        c.fill = FILL["scen"] if col in COLS else FILL["colhead"]
    labels = {row["key"]: row for row in ROWS if row.get("key")}
    wrap = Alignment(wrap_text=True, vertical="top")
    r = 4
    for key, label, bench, geo, urls, date, conf in SOURCES:
        r += 1
        src_row = labels[key]
        ws[f"B{r}"] = label or src_row["label"]
        ws[f"C{r}"] = src_row["unit"] if src_row["unit"] is not None else UNIT[src_row["fmt"]]
        for sc, col in enumerate(COLS):
            c = ws[f"{col}{r}"]
            c.value = f"={s1ref(key, sc, rowmap)}"
            c.number_format = FMT[src_row["fmt"]]
            c.font = font(color="008000")
            c.alignment = Alignment(horizontal="right", vertical="top")
        ws[f"G{r}"], ws[f"H{r}"], ws[f"M{r}"], ws[f"N{r}"] = bench, geo, date, conf
        for col in "BCGHMN":
            ws[f"{col}{r}"].alignment = wrap
            ws[f"{col}{r}"].font = font()
        if len(urls) > 4:
            raise ValueError(f"too many sources for {key}")
        for col, url in zip("IJKL", urls):
            c = ws[f"{col}{r}"]
            c.value = source_label(url)
            c.hyperlink = url
            c.font = font(color="0563C1", underline="single")
            c.alignment = wrap
        if not urls:
            ws[f"I{r}"] = "—"
            ws[f"I{r}"].font = font()
        ws.row_dimensions[r].height = max(28, 13 * (len(bench) // 58 + 1) + 4)
    ws.freeze_panes = "C5"
    return ws


def main():
    out, map_out = Path(sys.argv[1]), Path(sys.argv[2])
    rowmap = layout()
    wb = Workbook()
    build_sheet1(wb, rowmap)
    build_sheet2(wb, rowmap)
    build_sheet3(wb)
    build_sheet4(wb, rowmap)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    wb.save(out)
    s2 = dict(blocks=S2_BLOCKS, m0=S2_M0, months=S2_MONTHS, total=S2_TOTAL, chk_in=S2_CHK_IN, chk_out=S2_CHK_OUT,
              gap=S2_GAP, start=S2_START)
    map_out.write_text(json.dumps(dict(rowmap=rowmap, cols=COLS, s2=s2), ensure_ascii=False, indent=1))
    print(f"ok: {out} rows={len(ROWS)} keys={len(rowmap)}")


if __name__ == "__main__":
    main()
