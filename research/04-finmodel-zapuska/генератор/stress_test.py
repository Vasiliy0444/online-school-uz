#!/usr/bin/env python3
"""Стресс-тесты: меняем вводные прямо в книге, пересчитываем LibreOffice, сверяем с двойником."""
import copy, json, subprocess, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, str(Path(__file__).parent))
from model_defaults import DEF
from twin import compute, scenario
from verify_model import verify

RECALC = "/Users/Bel-a/Library/Application Support/Claude/local-agent-mode-sessions/skills-plugin/62ade98b-c20d-435d-bd1e-b573f3f7bff9/c9850139-8eac-49bf-9bdb-4c182fbb50f1/skills/xlsx/scripts/recalc.py"
meta = json.loads(Path("rowmap.json").read_text())
rowmap, cols = meta["rowmap"], meta["cols"]

CASES = {
    "budget_x2": {"budget": (10000, 10000, 10000)},
    "budget_0": {"budget": (0, 0, 0)},
    "no_own_installment": {"pay_full": (0.65, 0.70, 0.75)},
    "term_1": {"own_term": (1, 1, 1)},
    "gap6_term12": {"gap": (6, 6, 6), "own_term": (12, 12, 12)},
    "channel_start_3000_upsell": {"ch_start": (3000, 3000, 3000), "upsell_rate": (0.05, 0.1, 0.15), "w_budget": (500, 500, 500)},
    "zero_cpm_zero_wcpl": {"cpm": (0, 0, 0), "w_cpl": (0, 0, 0), "w_budget": (300, 300, 300)},
}

def run_case(name, overrides):
    d = copy.deepcopy(DEF)
    d.update(overrides)
    wb = load_workbook("model.xlsx")
    ws = wb["Декомпозиция цикла"]
    for key, vals in overrides.items():
        for sc, col in enumerate(cols):
            ws[f"{col}{rowmap[key]}"] = vals[sc]
    out = f"stress_{name}.xlsx"
    wb.save(out)
    res = json.loads(subprocess.run([sys.executable, RECALC, out, "90"], capture_output=True, text=True).stdout)
    errs = res.get("total_errors", "?")
    problems = verify(out, "rowmap.json", d, quiet=True)
    return d, errs, problems

base = {sc: compute(scenario(DEF, sc)) for sc in range(3)}
for name, ov in CASES.items():
    d, errs, problems = run_case(name, ov)
    r = {sc: compute(scenario(d, sc)) for sc in range(3)}
    extra = ""
    if name == "budget_x2":
        extra = " | reg ratio " + ", ".join(f"{r[sc]['reg']/base[sc]['reg']:.3f}" for sc in range(3)) + \
                " | cold sales ratio " + ", ".join(f"{r[sc]['c_sales']/base[sc]['c_sales']:.3f}" for sc in range(3))
    if name == "no_own_installment":
        extra = " | tail " + ", ".join(f"{r[sc]['cash_tail_total']:.2f}" for sc in range(3))
    if name == "term_1":
        extra = " | tail " + ", ".join(f"{r[sc]['cash_tail_total']:.2f}" for sc in range(3))
    if name == "budget_0":
        extra = " | reg " + ", ".join(f"{r[sc]['reg']:.0f}" for sc in range(3)) + " profit " + ", ".join(f"{r[sc]['profit_pre']:,.0f}" for sc in range(3))
    print(f"{name}: formula errors={errs} twin mismatches={problems}{extra}")
