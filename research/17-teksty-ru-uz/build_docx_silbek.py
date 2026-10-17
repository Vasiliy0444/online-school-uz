"""Docx для Силбека: тексты под постами, русский | узбекский | пустая колонка для его варианта.

Источники — pools_ru.json и pools_uz_final.json. Формат как у docx креативов
(16-kreativy-meta/docx_for_native.py), который Силбек уже возвращал 15.10.

python3 build_docx_silbek.py
"""
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "для-носителя" / "Тексты под постами RU-UZ — для Силбека.docx"
DESK_OUT = Path.home() / "Desktop" / OUT.name
GREY = RGBColor(0x6B, 0x6B, 0x6B)
COL_W = (Cm(8.6), Cm(8.6), Cm(8.5))
LENGTH = {"short": "Короткий", "mid": "Средний", "long": "Лонгрид"}

INTRO = [
    "Это рекламные тексты под постами в Instagram и Facebook для Узбекистана "
    "(подпись под картинкой или видео). Продукт — бесплатный мини-курс в Telegram «AI bilan 0 dan ilk $1000 gacha».",
    "9 групп креативов, на каждую — три текста: короткий, средний и длинный. Всего 27 текстов.",
    "Слева — русский текст, его мы согласовали. В середине — узбекский перевод, сделанный с помощью ИИ. "
    "Справа — пустая колонка для вашего варианта.",
]
ASKS = [
    "Если узбекский звучит естественно и смысл совпадает с русским — напишите в правой колонке «ок».",
    "Если что-то не так — напишите в правой колонке весь правильный текст и выделите жёлтым то, что поменяли.",
    "Латиница; o‘ и g‘ — с апострофом ‘. Обращение — на «siz».",
    "Цифры, суммы, порядок абзацев и эмодзи (⏳ 👆 👉 ✅ 1️⃣) не меняем. Если число кажется странным — напишите об этом.",
    "Не добавляйте обещаний вроде kafolat, tez boyish, oson pul, passiv daromad — такую рекламу Meta отклоняет.",
    "Кнопка под рекламой — «Batafsil» (Подробнее). Если узбеки так не говорят, предложите, как лучше.",
]


def shade(cell, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def fill_cell(cell, text: str, *, size=10, bold=False, color=None):
    """Абзацы текста — отдельные абзацы в ячейке."""
    cell.paragraphs[0].text = ""
    for i, part in enumerate(text.split("\n")):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        run = para.add_run(part)
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color


def clean(text: str) -> str:
    return "\n".join(line for line in text.split("\n") if line.strip())


def setup(doc):
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.6)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        if name != "Normal":
            st.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(15)


def pool_table(doc, ru: dict, uz: dict):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    head = table.rows[0]
    repeat_header(head)
    titles = ("Русский (согласовано)", "Узбекский (перевод ИИ)", "Вариант Силбека")
    for cell, title in zip(head.cells, titles):
        fill_cell(cell, title, bold=True)
        shade(cell, "EDEDED")
    for key, label in LENGTH.items():
        sub = table.add_row()
        merged = sub.cells[0].merge(sub.cells[2])
        fill_cell(merged, f"{label} · {len(ru['texts'][key])} знаков", size=9, bold=True, color=GREY)
        shade(merged, "F7F7F7")
        cells = table.add_row().cells
        fill_cell(cells[0], clean(ru["texts"][key]))
        fill_cell(cells[1], clean(uz["texts"][key]))
        fill_cell(cells[2], "")
        shade(cells[2], "FFFDF0")
    for row in table.rows:
        for cell, w in zip(row.cells, COL_W):
            cell.width = w


def main():
    ru_pools = json.loads((ROOT / "pools_ru.json").read_text())["pools"]
    uz_pools = {p["n"]: p for p in json.loads((ROOT / "pools_uz_final.json").read_text())["pools"]}
    doc = Document()
    setup(doc)
    doc.add_heading("Тексты под постами: русский и узбекский", level=0)
    for line in INTRO:
        doc.add_paragraph(line)
    doc.add_paragraph("Что нужно:").runs[0].bold = True
    for line in ASKS:
        doc.add_paragraph(line, style="List Bullet")

    for ru in ru_pools:
        uz = uz_pools[ru["n"]]
        if set(ru["texts"]) != set(uz["texts"]):
            raise KeyError(f"пул {ru['n']}: длины в RU и UZ не совпадают")
        h = doc.add_heading(f"Группа {ru['n']}. {ru['title']} / {uz['title_uz']}", level=1)
        h.paragraph_format.page_break_before = True
        h.paragraph_format.keep_with_next = True
        p = doc.add_paragraph()
        r = p.add_run(f"Креативы: {ru['creatives']}")
        r.font.size, r.font.color.rgb = Pt(9), GREY
        p.paragraph_format.keep_with_next = True
        pool_table(doc, ru, uz)

    doc.add_heading("Общие замечания", level=1).paragraph_format.page_break_before = True
    doc.add_paragraph("Если есть что-то общее для всех текстов — напишите здесь.")
    for _ in range(6):
        doc.add_paragraph("")

    doc.save(OUT)
    shutil.copy(OUT, DESK_OUT)
    print(f"{len(ru_pools)} групп, {len(ru_pools) * 3} текстов → {OUT}\nкопия: {DESK_OUT}")


if __name__ == "__main__":
    main()
