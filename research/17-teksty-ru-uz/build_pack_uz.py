#!/usr/bin/env python3
"""Собирает папку «для-носителя»: xlsx на проверку узбекским ассистентом + README.

    python3 build_pack_uz.py pools_uz_final.json

В xlsx три листа: «Как проверять», «Тексты» (русский рядом с узбекским и колонки для правок),
«Вопросы» (спорные места, где нужен именно живой носитель).
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "для-носителя"
LABELS = {"short": "Короткий", "mid": "Средний", "long": "Лонгрид"}
STATUSES = "ок,мелкая правка,переписать"

HOW_TO = [
    "Что это: рекламные тексты под постами в Instagram и Facebook для Узбекистана. Продукт: бесплатный мини-курс в Telegram «AI bilan 0 dan ilk $1000 gacha».",
    "Перевод сделал ИИ, живой носитель его ещё не читал. Ваша задача: убедиться, что тексты звучат так, как их написал бы узбек, а не перевод с русского.",
    "Лист «Тексты»: слева русский оригинал (только для сверки смысла), рядом узбекский на проверку. Заполните три колонки справа: «Что не так», «Ваш вариант», «Статус».",
    "Лист «Вопросы»: места, где без носителя не решить (по убыванию важности для пулов 1–6, потом для пулов 7–9). Ответьте в колонке «Ответ».",
    "Что НЕ менять: числа и суммы ($1000, $300, $100, 3 mln so‘m, 100 ta joy, 26 ta), порядок абзацев и эмодзи (⏳ 👆 👉 ✅ 1️⃣…5️⃣). Если число кажется странным, напишите в «Что не так», но не правьте.",
    "Орфография: o‘ и g‘ пишем со знаком ‘ (как в этом файле), обращение везде «siz», латиница.",
    "Не добавляйте новых обещаний (kafolat, tez boyish, oson pul, passiv daromad): рекламу такого рода Meta отклоняет.",
    "Если фраза нормальная, поставьте «ок» и ничего не пишите. Правка нужна там, где носитель так не скажет, или звучит как русская калька, или двусмысленно.",
]

QUESTIONS = [
    ("Кнопка. Как в узбекском интерфейсе Instagram и Facebook подписана кнопка «Подробнее» (Learn more)? В текстах стоит «Batafsil» (17 мест). Оставить так, или заменить везде на «Ro‘yxatdan o‘tish uchun yuqoridagi tugmani bosing» без названия кнопки?", "Все тексты, последний абзац"),
    ("«AI» в роли подлежащего: «Asosiy ishni AI qiladi», «AI bajaradi», «Bu rolikni AI yaratdi». Звучит естественно? Или лучше «sun’iy intellekt», «neyrotarmoq»?", "P2 short/mid/long, P3, P4, P6"),
    ("«yaqin atrofingizdagi 3 ta kompaniya» (русское «3 компании в своём районе»). Естественно? Как сказать лучше: «mahallangizdagi», «tumaningizdagi», «shahringizdagi»?", "long-тексты P1, P3, P4, P5, P6"),
    ("Хук пула 2: «Vaqtingizni arzimagan pulga sotyapsizmi?» («Продаёшь своё время за копейки?»). Понятно и звучит ли? Или «arzon-garovga»?", "P2 short/mid/long"),
    ("«Ishlagan soat(ingiz) uchun to‘lashsa» (русское «когда платят за часы»). Естественно ли?", "P2 long, P3 mid, P3 long"),
    ("Порядок слов: «Sayt, Telegram-bot, dizayn va Reels uchun biznes pul to‘laydi». Не читается ли как «бизнес-сайт»? Как лучше?", "P4 mid, P4 long"),
    ("«Talaba bo‘lsangiz va internetda hali bir marta ham pul ishlamagan bo‘lsangiz…» («если ты учишься»). Не слишком узко (только вуз и колледж)? Школьники и ученики курсов не попадают. Нужна ли более широкая формулировка?", "P1 mid"),
    ("«Aksiya davom etar ekan, kurs to‘liq bepul» («пока действует предложение»). Естественно ли «ekan» в значении «пока»?", "P1 mid"),
    ("Кнопка или ссылка: «yuqoridagi tugmani bosing» и «yuqoridagi havola orqali». В Instagram над подписью стоит кнопка, как лучше называть?", "P1 short, все призывы 👆"),
    ("Суммы и лимит: «3 mln o‘rniga 0 so‘m», «100 ta joydan 26 tasi qoldi», «Jami 100 ta joy, 74 tasi allaqachon band». Так принято писать по-узбекски?", "почти все тексты"),
    ("Тон срочности: ⏳ «kun oxirigacha», «ulguring». Не слишком ли давит, не звучит ли как обман?", "почти все тексты"),
    ("Хук пула 3: вопрос от первого лица в кавычках «Kun bo‘yi ishlayman, lekin pulim yetmayapti. Qanday qilib ko‘proq pul topsam bo‘ladi?». Естественно как «вопрос подписчика»?", "P3 short/mid/long"),
    ("Пул 8, хук: «Ko‘p odam chetga ishlagani ketadi» («Многие уезжают на заработки»). Звучит ли нейтрально, без жалости и оценки? «chetga» или «chet elga»? Если менять, то вместе с каруселью 2-4.", "P8 short/mid/long"),
    ("Пул 8: «Bilak kuchi bilan emas, bilim bilan ishlaysiz» («работаешь головой, а не спиной»). Не обидно ли тем, у кого родные на заработках? Естественно ли «bilim bilan ishlash»?", "P8 mid, long"),
    ("Пул 9, хук: «Yana «AI orqali $1000» va’dasimi? Quruq gapga ishonmang». Не бьёт ли «va’da» по нашему же курсу про $1000? Не слишком ли резко «quruq gap»? Альтернатива: «Yana «AI orqali $1000»mi?».", "P9 short/mid/long"),
    ("Пул 8: «Uyda ham mijozlar bor» («дома заказчики тоже есть»). Читается как «на родине» или как «в доме»? Альтернатива: «O‘zimizda ham mijozlar bor».", "P8 mid, long"),
    ("Пул 7: «pulni esa mijoz topadiganlarga to‘lashadi» («платят тем, кто находит заказчиков»). Естественно ли?", "P7"),
    ("Пул 7, хук: «Mijoz topishni bilmaysizmi? Xizmatingizni arzonga sotasizmi?» Не звучит ли упрёком? «bilmaysizmi» или «bilmayapsizmi»? Менять только вместе с картинкой I3-3.", "P7 short/mid/long"),
    ("Пул 8 long: «— va oila yoningizda qoladi» рядом с «oila esa uzoqda qoladi». Естественно ли?", "P8 long"),
    ("Мелочи: «mo‘ljal shu», «beqaror», «munosib narx» (пул 7), «Kod yozmadim» (пул 9).", "P7, P9"),
]


def style_header(ws, row=1):
    for cell in ws[row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F5D50")
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def build_xlsx(ru_pools, uz_pools, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Как проверять"
    ws.column_dimensions["A"].width = 140
    for i, line in enumerate(HOW_TO, start=1):
        ws.cell(row=i, column=1, value=line).alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Тексты")
    ws.append(["ID", "Пул", "Длина", "Русский (для сверки)", "Узбекский (проверить)", "Что не так", "Ваш вариант", "Статус"])
    style_header(ws)
    for ru_pool, uz_pool in zip(ru_pools, uz_pools):
        for key, label in LABELS.items():
            ws.append([
                f"P{ru_pool['n']}-{key}",
                f"{ru_pool['n']}. {ru_pool['title']}",
                label,
                ru_pool["texts"][key],
                uz_pool["texts"][key],
                "", "", "",
            ])
    for col, width in zip("ABCDEFGH", (10, 30, 11, 55, 55, 35, 55, 16)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    validation = DataValidation(type="list", formula1=f'"{STATUSES}"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"H2:H{ws.max_row}")

    ws = wb.create_sheet("Вопросы")
    ws.append(["№", "Вопрос", "Где в текстах", "Ответ проверяющего"])
    style_header(ws)
    for n, (question, where) in enumerate(QUESTIONS, start=1):
        ws.append([n, question, where, ""])
    for col, width in zip("ABCD", (5, 90, 30, 60)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def build_readme(path):
    text = "# Для проверяющего: тексты под постами, узбекский\n\n"
    text += "\n".join(f"- {line}" for line in HOW_TO)
    text += "\n\nФайл с текстами: `Тексты на проверку — UZ.xlsx`. Верните его заполненным.\n"
    path.write_text(text, encoding="utf-8")


def main(uz_path):
    ru = json.loads((HERE / "pools_ru.json").read_text(encoding="utf-8"))["pools"]
    uz = json.loads(Path(uz_path).read_text(encoding="utf-8"))["pools"]
    OUT_DIR.mkdir(exist_ok=True)
    build_xlsx(ru, uz, OUT_DIR / "Тексты на проверку — UZ.xlsx")
    build_readme(OUT_DIR / "README.md")
    print(f"готово: {OUT_DIR}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else HERE / "pools_uz_v2.json")
