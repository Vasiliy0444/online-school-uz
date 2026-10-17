#!/usr/bin/env python3
"""Сверяет пересчитанный шаблон с Python-расчётом. Использование: python3 verify_template.py <xlsx> <addr.json>"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from template_defaults import AI_SCEN, COLD, S3, S5, S6, TRAFFIC, WARM  # noqa: E402
from template_twin import compute  # noqa: E402

COLD_KEYS = ["reg", "part", "peak", "orders", "orders_sum", "autopay", "autopay_sum", "prepay", "op_leads", "op_bk_plan",
             "op_or_plan", "cv_web", "web_plan", "op_paid", "op_sum", "traffic", "roi", "plan"]
WARM_KEYS = ["cold_reg", "ch_from_cold", "ch_buyers", "base", "reg", "part", "autopay", "bk", "orders", "op_leads",
             "op_bk_plan", "op_or_plan", "cv_dozhim", "dozhim_plan", "op_total", "op_sum", "total", "plan"]


def verify(path, addr_path, params=None, quiet=False):
    p = params or dict(cold=COLD, warm=WARM, ai_scen=AI_SCEN, s3=S3, s5=S5, s6=S6, traffic=TRAFFIC)
    tw = compute(**p)
    addr = json.loads(Path(addr_path).read_text())
    wb = load_workbook(path, data_only=True)
    problems = 0

    def cell(key):
        sh, col, row = addr[key]
        return wb[sh][f"{col}{row}"].value

    def check(key, exp, tol=0.01):
        nonlocal problems
        got = cell(key)
        if got is None or isinstance(got, str) or abs(float(got) - exp) > max(tol, abs(exp) * 1e-6):
            problems += 1
            print(f"MISMATCH {key}: book={got} twin={exp}")

    for bid in ("B_base", "B_min", "B_mid", "B_max"):
        for k in COLD_KEYS:
            check(f"{bid}:{k}", tw[bid][k])
    for bid in ("A_base", "A_min", "A_opt", "A_max"):
        for k in WARM_KEYS:
            check(f"{bid}:{k}", tw[bid][k])
    for k in ("op_sum", "plan", "total", "op_total"):
        check(f"C_s1:{k}", tw["A_base"][k])
    for k in ("op_paid", "cv_op", "op_sum", "plan", "roi"):
        check(f"C_s2:{k}", tw["C_s2"][k])
    for k in ("noact", "noshow", "dialogs", "ai_paid", "hot", "m_paid", "total", "plan"):
        check(f"C_s3:{k}", tw["C_s3"][k])
    for k in ("op_plan", "launch_plan"):
        check(f"C_s4:{k}", tw["C_s4"][k])
    for k in ("cold_rev", "warm_rev", "rev", "cold_now", "cold_month", "warm_now", "warm_month", "now", "later",
              "loss_acq", "loss_fee", "loss_def", "loss_ref", "net"):
        check(f"C_s5:{k}", tw["C_s5"][k])
    for k in ("ads", "months", "cold_students", "warm_students", "cold_mgr_rev", "warm_mgr_rev", "cost_ai", "cost_mgr",
              "cost_team", "cost_services", "cost_cur", "cost_tax", "costs", "profit", "profit_rec", "owner", "romi",
              "start_cash", "cold_var", "warm_var", "fixed_month"):
        check(f"C_s6:{k}", tw["C_s6"][k])
    check("C_cash:low", tw["cash"]["low"])
    # касса по месяцам и сверки
    sh, col, low_row = addr["C_cash:low"]
    ws = wb[sh]
    total_row = low_row - 4
    first_row = total_row - 18
    for mo, inflow, outflow, net, cum in tw["cash"]["flows"]:
        r = first_row + mo - 1
        for c, exp in (("L", inflow), ("M", outflow), ("O", cum)):
            got = ws[f"{c}{r}"].value
            if got is None or abs(float(got) - exp) > 0.01:
                problems += 1
                print(f"CASH {c}{r} m{mo}: book={got} twin={exp}")
    for c, r in (("L", total_row + 2), ("M", total_row + 3)):
        v = ws[f"{c}{r}"].value
        if v is None or abs(float(v)) > 0.01:
            problems += 1
            print(f"RECON {c}{r}: {v}")
    # «Трафик» и «От куда?»
    tr = wb["Трафик и лидмагнитная воронка"]
    if abs(float(tr["D6"].value) - tw["traffic"]["target_reg"]) > 0.01 or abs(float(tr["C9"].value) - tw["traffic"]["budget_total"]) > 0.01:
        problems += 1
        print("TRAFFIC mismatch", tr["D6"].value, tr["C9"].value)
    src = wb["Модель продаж- через вебинармар"]
    if abs(float(src["I8"].value) - tw["traffic"]["target_reg"]) > 0.01:
        problems += 1
        print("SOURCES total mismatch", src["I8"].value)
    if not quiet:
        s6, s5, s4 = tw["C_s6"], tw["C_s5"], tw["C_s4"]
        print(f"цикл: выручка ${s5['rev']:,.0f} | чистые деньги ${s5['net']:,.0f} (сразу ${s5['now']:,.0f}, позже ${s5['later']:,.0f}) "
              f"| расходы ${s6['costs']:,.0f} | прибыль ${s6['profit']:,.0f} | собственнику ${s6['owner']:,.0f} | "
              f"ROMI {s6['romi']:.0%} | касса min ${tw['cash']['low']:,.0f}")
        for bid in ("B_min", "B_mid", "B_max"):
            print(f"  {bid}: план продаж ${tw[bid]['plan']:,.0f}, окупаемость {tw[bid]['roi']:.0%}")
        for bid in ("A_min", "A_opt", "A_max"):
            print(f"  {bid}: база канала {tw[bid]['base']:.0f}, регистраций {tw[bid]['reg']:.0f}, план ${tw[bid]['plan']:,.0f}")
    print("PROBLEMS:", problems)
    return problems


if __name__ == "__main__":
    sys.exit(1 if verify(sys.argv[1], sys.argv[2]) else 0)
