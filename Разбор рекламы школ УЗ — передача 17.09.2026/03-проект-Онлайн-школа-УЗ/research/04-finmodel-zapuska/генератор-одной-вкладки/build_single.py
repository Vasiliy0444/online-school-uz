#!/usr/bin/env python3
"""Одна вкладка «Модель продаж - гибридная» на шаблоне заказчика.

Слева — этап 1: запуск через вебинар (общий план, «От куда?», декомпозиция, 3 сценария, подписчики ТГ-канала).
Справа — этап 2: запуск через ТГ-канал, марафон 2 дня (3 сценария).
Итог двух этапов — по средним сценариям: деньги за 2 запуска и окупаемость.

Использование: python3 build_single.py <out.xlsx> <addr.json>
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import column_index_from_string as cidx
from openpyxl.utils import get_column_letter as cl

sys.path.insert(0, str(Path(__file__).parent))
from single_model import AI, S1, S2, TOT, blended_rates  # noqa: E402

TEMPLATE = "/Users/Bel-a/Downloads/декомпозиция Узбекситан .xlsx"
SHEET = "Модель продаж - гибридная"

THIN = Side(style="thin", color="FF000000")
NO = Side()
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
C_TITLE, C_NOTE, C_WHITE, C_TOT = "FF00FFFF", "FFFFF2CC", "FFFFFFFF", "FFB6D7A8"
C_SEC_A, C_SEC_B, C_IN_A, C_IN_B = "FFFFE599", "FFFFF2CC", "FFCFE2F3", "FFC9DAF8"
FMT = {"int": "#,##0", "pct": "0.0%", "pct2": "0.00%", "usd": '"$"#,##0;[Red]-"$"#,##0', "usd2": '"$"#,##0.00',
       "roi": "0%", "mon": "0", "gen": "General"}

ADDR = {}
TOKEN = re.compile(r"\{(?:(\w+):)?(\w+)\}")


def brd(l=False, r=False, t=False, b=False):  # noqa: E741
    return Border(left=THIN if l else NO, right=THIN if r else NO, top=THIN if t else NO, bottom=THIN if b else NO)


def mfont(size=11, bold=False, italic=False):
    return Font(name="Montserrat", size=size, bold=bold, italic=italic, color=Color(theme=1))


def solid(rgb):
    return PatternFill("solid", fgColor=rgb)


def R(key, label, kind, fmt, v=None, note=None):
    return dict(key=key, label=label, kind=kind, fmt=fmt, v=v, note=note)


class Block:
    def __init__(self, bid, sections, sec_col, lbl_col, val_col, title, title_row, title_fill=C_TITLE,
                 sec_fill=C_SEC_A, in_fill=C_IN_A, title_height=None):
        self.__dict__.update(bid=bid, sections=sections, sec_col=sec_col, lbl_col=lbl_col, val_col=val_col,
                             title=title, title_row=title_row, title_fill=title_fill, sec_fill=sec_fill,
                             in_fill=in_fill, title_height=title_height)
        r = title_row + 1
        for _label, rows in sections:
            for row in rows:
                row["r"] = r
                if (bid, row["key"]) in ADDR:
                    raise ValueError(f"duplicate {bid}:{row['key']}")
                ADDR[(bid, row["key"])] = (val_col, r)
                r += 1
        self.last_row = r - 1


def resolve(expr, bid):
    def rep(m):
        col, row = ADDR[(m.group(1) or bid, m.group(2))]
        return f"{col}{row}"
    return TOKEN.sub(rep, expr)


def merge_box(ws, c1, r1, c2, r2, value, font, fill, height=None, borders=True):
    ws.merge_cells(f"{c1}{r1}:{c2}{r2}")
    top = ws[f"{c1}{r1}"]
    top.value, top.font, top.fill, top.alignment = value, font, fill, CENTER
    if borders:
        for rr in range(r1, r2 + 1):
            for ci in range(cidx(c1), cidx(c2) + 1):
                ws[f"{cl(ci)}{rr}"].border = brd(l=ci == cidx(c1), r=ci == cidx(c2), t=rr == r1, b=rr == r2)
    if height:
        ws.row_dimensions[r1].height = height


def write_block(ws, blk):
    merge_box(ws, blk.sec_col, blk.title_row, blk.val_col, blk.title_row, blk.title, mfont(11, True),
              solid(blk.title_fill), blk.title_height)
    for label, rows in blk.sections:
        r1, r2 = rows[0]["r"], rows[-1]["r"]
        if r2 > r1:
            ws.merge_cells(f"{blk.sec_col}{r1}:{blk.sec_col}{r2}")
        c = ws[f"{blk.sec_col}{r1}"]
        c.value, c.font, c.fill, c.alignment = label, Font(name="Arial", bold=True, color=Color(theme=1)), solid(blk.sec_fill), CENTER
        for rr in range(r1, r2 + 1):
            ws[f"{blk.sec_col}{rr}"].border = brd(l=True, r=True, t=rr == r1, b=rr == r2)
        for row in rows:
            r, kind = row["r"], row["kind"]
            fill = {"in": blk.in_fill, "f": C_WHITE, "tot": C_TOT, "tot_in": C_TOT}[kind]
            lc, vc = ws[f"{blk.lbl_col}{r}"], ws[f"{blk.val_col}{r}"]
            lc.value = row["label"]
            lc.font, lc.fill, lc.alignment, lc.border = mfont(11, True), solid(fill), CENTER, brd(l=True, r=True, b=True)
            v = row["v"]
            vc.value = resolve(v, blk.bid) if isinstance(v, str) and v.startswith("=") else v
            vc.font, vc.fill, vc.alignment = mfont(11, kind in ("tot", "tot_in")), solid(fill), CENTER
            vc.border, vc.number_format = brd(r=True, b=True), FMT[row["fmt"]]
            if row.get("note"):
                cm = Comment(row["note"], "Claude")
                cm.width, cm.height = 330, 150
                vc.comment = cm


NOTE = {
    "cpl": "Лид в образовании УЗ — 2,3–3,6 $ (https://aksanov.digital/blog/meta-ads/analiz-kampaniy); в сезон июль–сентябрь "
           "10–12 $ (https://t.me/s/targetologi_uz); база брифа рынка — 6 $.",
    "att": "Доходимость до вебинара: в среднем 15–20% (https://www.sostav.ru/blogs/287076/96331); 20–30% — норма, 50% — хорошо "
           "(https://vc.ru/id972483/798734-dohodimost-do-vebinara-kakoi-procent-optimalnyi).",
    "conv": "Со вебинара покупают 5–20% зрителей (skillspace.ru, РФ); в брифе рынка заявки — 12% участников, оплачивают 40%. "
            "Доли автооплат и предоплат — оценка.",
    "op_bk": "Заказ → оплата 66,8% (https://getcourse.ru/magazine/analitika_2025); в примере финмодели школы — 46%.",
    "op_or": "Звонок → оплата 50–55% по гайду (не проверено); база брифа рынка — 40%.",
    "dozhim1": "Дожим в боте всех, кто зарегистрировался и не купил: ИИ-продавец пишет текстом, горячих закрывает менеджер. "
               "Считается от «регистраций − заказы − автооплаты − предоплаты». Ставка — оценка: ответили 20–30% "
               "бывших на вебинаре и 6–10% не дошедших, из диалогов 1–2% платят сами и 6–10% уходят менеджеру.",
    "dozhim2": "Дожим в боте тех, кто зарегистрировался на марафон и не купил: ИИ-продавец + менеджер. Ставка — оценка.",
    "check": "Чек 300 $ — по ответу Васи. У конкурентов в УЗ: фриланс 170–595 $, ИИ-академии 255–510 $; база брифа — 212 $.",
    "peak": "Из шаблона: пик = 60% участников. Вынесено из формулы в ячейку.",
    "tg_sub": "Сколько зарегистрировавшихся подпишутся на ТГ-канал школы. Оценка — данных по УЗ нет.",
    "churn": "Отписки 8–12% подписчиков (https://vc.ru/telegram/1425019-kak-poschitat-okupaemost-telegram-kanala, РФ).",
    "warm_att": "Холодный марафон: день 1 — 25%, день 2 — 14% регистраций; тёплый: день 1 — 68%, на день 2 вернулись 72% (vc.ru, РФ).",
    "est": "Оценка — данных по Узбекистану нет.",
    "acq": "Payme / Click / Uzum: 0,8–2% (https://uzneo.uz/blog/payme-click-uzum-podklyuchit-oplatu).",
    "partner": "Рассрочка 0% на 6–12 мес через Uzum и Alif — стандарт рынка УЗ (бриф рынка). Доля 25% — оценка: 30% реальна только при двух провайдерах и 0% на 12 мес. Стартовый лимит Uzum Nasiya — 800 тыс. сум (https://t.me/uzumnasiyauz/704), курс дороже. Подробно — research/06-intervyu-praktika/РАССРОЧКИ-ДЛЯ-ШКОЛ.md.",
    "fee": "Ставки Uzum и Alif для школ не публикуются — запросить КП (задача CURSOR-q1z0). Оценка по стоимости денег: 3 мес ~7%, 6 мес ~12%, 12 мес ~20%. Проверка: при наценке покупателю 27–44% за 12 мес эквивалентная скидка 21–31% (https://www.gazeta.uz/ru/2025/07/02/nasiya-iphone/).",
    "refund": "Возвраты 5% — допущение брифа рынка.",
    "ai_setup": "Разработка ИИ-бота в УЗ — 5–25 млн сум; ИИ и хостинг — 0,3–3 млн сум в месяц (https://innosoft.uz/blog/ai-narxi-2026).",
    "tax": "ООО — налог с оборота 4%. Порог с 01.06.2026 — 12 000 БРВ ≈ 5 млрд сум (https://lex.uz/ru/docs/8237115); в тексте НК пока 1 млрд — уточнить у юриста.",
    "curator": "Кураторы 10–25 $ на ученика — допущение брифа рынка.",
    "m_pct": "Затраты на продажи 10% выручки — допущение брифа рынка.",
}


def stage1_sections(i, mirror=False, plan=False):
    """i: 0 — минимум, 1 — средний, 2 — максимум. plan — общий план (вводные); mirror — средний сценарий = общий план."""
    cv1 = CV1[i]

    def val(key, v):
        return "={S1_plan:%s}" % key if mirror else v

    def kind_in():
        return "f" if mirror else "in"
    n = (lambda k: NOTE[k]) if plan else (lambda k: None)
    reg_v = "={SRC:total}" if plan else ("={S1_plan:reg}" if mirror else "=IF({cpl}>0,{traffic}/{cpl},0)")
    return [
        ("МАРКЕТИНГ И ВОРОНКА", [
            R("reg", "Регистраций:", "f", "int", reg_v,
              "Сумма таблицы «От куда?» справа; таргет = вложили в трафик ÷ цена регистрации." if plan else None),
            R("cpl", "Цена регистрации (таргет)", kind_in(), "usd2", val("cpl", S1["cpl"][i]), n("cpl")),
            R("att", "Доходимость:", kind_in(), "pct", val("att", S1["att"][i]), n("att")),
            R("part", "Участников на вебинаре:", "f", "int", "={reg}*{att}"),
            R("peak", "Пик:", "f", "int", "={part}*{peak_pct}"),
            R("peak_pct", "Пик, % от участников", kind_in(), "pct", val("peak_pct", S1["peak_pct"][i]), n("peak")),
            R("orders", "Оставили заказ:", "f", "int", "={part}*{cv_order}"),
            R("orders_sum", "Сумма заказов:", "f", "usd", "={orders}*{check}"),
            R("cv_order", "Конверсия из участника в заказ", kind_in(), "pct", val("cv_order", S1["cv_order"][i]), n("conv")),
            R("autopay", "Автооплат шт", "f", "int", "={part}*{cv_autopay}"),
            R("autopay_sum", "Сумма автооплат", "f", "usd", "={autopay}*{check}"),
            R("cv_autopay", "Конверсия в автооплаты", kind_in(), "pct", val("cv_autopay", S1["cv_autopay"][i])),
            R("prepay", "Предоплат шт", "f", "int", "={part}*{cv_prepay}"),
            R("cv_prepay", "Конверсия в предоплату", kind_in(), "pct", val("cv_prepay", S1["cv_prepay"][i])),
        ]),
        ("ОТДЕЛ ПРОДАЖ", [
            R("op_leads", "Заявок на оп", "f", "int", "={orders}+{prepay}"),
            R("op_leads_sum", "Сумма заявок на оп", "f", "usd", "={op_leads}*{check}"),
            R("op_bk", "CV ОП - БРОНЬ", kind_in(), "pct", val("op_bk", S1["op_bk"][i]), n("op_bk")),
            R("op_bk_plan", "План закрытые брони", "f", "int", "={op_bk}*{prepay}"),
            R("op_or", "СV ОП - ЗАКАЗ", kind_in(), "pct", val("op_or", S1["op_or"][i]), n("op_or")),
            R("op_or_plan", "План закрытые заявки", "f", "int", "={op_or}*{orders}"),
            R("cv_web", "CV ОП - ВЕБИНАРНЫЕ: дожим в боте (ИИ-продавец + менеджер)", kind_in(), "pct2",
              val("cv_web", cv1), NOTE["dozhim1"] if (plan or not mirror) else None),
            R("web_plan", "План закрытия вебинарные", "f", "int", "={cv_web}*MAX(0,{reg}-{orders}-{autopay}-{prepay})"),
        ]),
        ("ИТОГИ", [
            R("op_paid", "Оплат оп", "f", "int", "={web_plan}+{op_or_plan}+{op_bk_plan}"),
            R("op_sum", "Сумма с оп", "f", "usd", "={op_paid}*{check}"),
            R("traffic", "Вложили в трафик", kind_in(), "usd", val("traffic", S1["traffic"][i])),
            R("roi", "Окупаемость трафика", "f", "roi", "=IF({traffic}>0,{plan}/{traffic},0)"),
            R("plan", "План продаж", "tot", "usd", "={op_sum}+{autopay_sum}"),
            R("check", "Ср. чек", kind_in(), "usd", val("check", S1["check"][i]), n("check")),
        ]),
        ("ТГ-КАНАЛ", [
            R("tg_sub", "Подписались на ТГ-канал, % регистраций", kind_in(), "pct", val("tg_sub", S1["tg_sub"][i]), n("tg_sub")),
            R("tg_sub_n", "Подписались на ТГ-канал, чел", "tot", "int", "={reg}*{tg_sub}"),
        ]),
    ]


def stage2_sections(i, s1bid):
    note = (lambda k: NOTE[k]) if i == 1 else (lambda k: None)
    return [
        ("ПРОГРЕВ", [
            R("from_s1", "Подписались на ТГ-канал в этапе 1", "f", "int", "={%s:tg_sub_n}" % s1bid),
            R("ch_start", "Подписчиков в канале было до этапа 1", "in", "int", S2["ch_start"][i],
              "Если у эксперта уже есть канал — впиши." if i == 1 else None),
            R("ch_org", "Органический прирост за 2 месяца", "in", "int", S2["ch_org"][i], note("est")),
            R("ch_churn", "Отписки за 2 месяца, %", "in", "pct", S2["ch_churn"][i], note("churn")),
            R("buyers", "Уже купили в этапе 1 — не целевые", "f", "int",
              "=({%s:op_paid}+{%s:autopay})*{%s:tg_sub}" % (s1bid, s1bid, s1bid)),
            R("base", "Средние охваты: база ТГ-канала к запуску", "f", "int",
              "=MAX(0,({ch_start}+{from_s1}+{ch_org})*(1-{ch_churn})-{buyers})"),
            R("reg", "Людей в анкете: регистраций на марафон", "f", "int", "={base}*{cv_reg}"),
            R("cv_reg", "CV охват - анкета: база канала → регистрация", "in", "pct", S2["cv_reg"][i], note("est")),
        ]),
        ("МАРКЕТИНГ ВОРОНКА", [
            R("part", "Дошли до продающего дня марафона", "f", "int", "={reg}*{att1}*{att2}"),
            R("att1", "Доходимость: день 1, % регистраций", "in", "pct", S2["att1"][i], note("warm_att")),
            R("att2", "Доходимость: день 2 (продажа), % от дня 1", "in", "pct", S2["att2"][i]),
            R("autopay", "Автооплаты", "f", "int", "={part}*{cv_autopay}"),
            R("cv_autopay", "СV автооплата - участники", "in", "pct", S2["cv_autopay"][i]),
            R("bk", "Брони", "f", "int", "={part}*{cv_bk}"),
            R("cv_bk", "СV бронь - участники", "in", "pct", S2["cv_bk"][i]),
            R("orders", "Заказы", "f", "int", "={part}*{cv_order}"),
            R("cv_order", "CV заказ - участники", "in", "pct", S2["cv_order"][i]),
            R("op_leads", "Заявок в работу на ОП", "f", "int", "={bk}+{orders}"),
        ]),
        ("ОТДЕЛ ПРОДАЖ", [
            R("op_bk", "CV ОП - БРОНЬ", "in", "pct", S2["op_bk"][i], note("op_bk")),
            R("op_bk_plan", "План закрытые брони", "f", "int", "={bk}*{op_bk}"),
            R("op_or", "СV ОП - ЗАКАЗ", "in", "pct", S2["op_or"][i], note("op_or")),
            R("op_or_plan", "План закрытые заявки", "f", "int", "={orders}*{op_or}"),
            R("cv_dozhim", "CV ОП - АНКЕТА: дожим в боте (ИИ-продавец + менеджер)", "in", "pct2", CV2[i], NOTE["dozhim2"]),
            R("dozhim_plan", "План закрытия анкеты: дожим тех, кто не купил", "f", "int",
              "=MAX(0,{reg}-{autopay}-{bk}-{orders})*{cv_dozhim}"),
        ]),
        ("ИТОГИ", [
            R("op_total", "Итого продаж ОП", "tot", "int", "={dozhim_plan}+{op_or_plan}+{op_bk_plan}"),
            R("op_sum", "План ОП сумма", "tot", "usd", "={op_total}*{check}"),
            R("autopay_total", "Итого автооплат", "tot", "int", "={autopay}"),
            R("total", "Итог общ", "tot", "int", "={op_total}+{autopay_total}"),
            R("check", "Средний чек", "tot_in", "usd", S2["check"][i], note("check")),
            R("plan", "План продаж", "tot", "usd", "={total}*{check}"),
        ]),
    ]


def total_sections():
    t = TOT
    first = "IF(ROUND({term},0)<=1,1,{first})"
    return [
        ("ПРОДАЖИ", [
            R("rev1", "План продаж этап 1 — вебинар (средний)", "f", "usd", "={S1_mid:plan}"),
            R("rev2", "План продаж этап 2 — ТГ-канал (средний)", "f", "usd", "={S2_mid:plan}"),
            R("op_plan", "ПЛАН ОП ПРОДАЖ", "f", "usd", "={S1_mid:op_sum}+{S2_mid:op_sum}"),
            R("launch", "ПЛАН ПРОДАЖ ЗАПУСК: 2 запуска за 2 месяца", "tot", "usd", "={rev1}+{rev2}"),
            R("students", "Продаж за 2 запуска, шт", "f", "int", "={S1_mid:op_paid}+{S1_mid:autopay}+{S2_mid:total}"),
        ]),
        ("ДЕНЬГИ", [
            R("full", "Полная оплата, % продаж", "in", "pct", t["full"]),
            R("partner", "Рассрочка Uzum Nasiya / Alif, % продаж", "in", "pct", t["partner"], NOTE["partner"]),
            R("own", "Своя рассрочка школы, % продаж (остаток)", "f", "pct", "=MAX(0,1-{full}-{partner})"),
            R("acq", "Эквайринг Payme / Click / Uzum, %", "in", "pct", t["acq"], NOTE["acq"]),
            R("fee", "Комиссия рассрочки Uzum / Alif для школы, %", "in", "pct", t["fee"], NOTE["fee"]),
            R("term", "Своя рассрочка: срок, мес", "in", "mon", t["term"]),
            R("first", "Своя рассрочка: первый платёж, % цены", "in", "pct", t["first"], NOTE["est"]),
            R("default", "Своя рассрочка: не доплатят, % остатка", "in", "pct", t["default"], NOTE["est"]),
            R("refund", "Возвраты, % выручки", "in", "pct", t["refund"], NOTE["refund"]),
            R("now", "Пришло денег сразу", "f", "usd",
              "={launch}*({full}*(1-{acq})+{partner}*(1-{fee})+{own}*%s*(1-{acq}))-{launch}*{refund}" % first),
            R("later", "Придёт позже частями по своей рассрочке", "f", "usd",
              "={launch}*{own}*(1-%s)*(1-{default})*(1-{acq})" % first),
            R("losses", "Съели эквайринг, комиссии, невыплаты и возвраты", "f", "usd", "={launch}-{now}-{later}"),
            R("net", "Чистые деньги за 2 запуска", "tot", "usd", "={now}+{later}"),
        ]),
        ("РАСХОДЫ", [
            R("traffic", "Вложили в трафик (этап 1, средний)", "f", "usd", "={S1_mid:traffic}"),
            R("ai_setup", "ИИ-продавец: разработка, разово", "in", "usd", t["ai_setup"], NOTE["ai_setup"]),
            R("ai_sub", "ИИ-продавец: ИИ и хостинг, $ в месяц", "in", "usd", t["ai_sub"], NOTE["ai_setup"]),
            R("salary", "Менеджер: оклад, $ в месяц", "in", "usd", t["salary"], NOTE["est"]),
            R("m_pct", "Менеджер: % с плана ОП продаж", "in", "pct", t["m_pct"], NOTE["m_pct"]),
            R("team", "Команда на 2 запуска: таргетолог, техспец, дизайнер, монтаж", "in", "usd", t["team"], NOTE["est"]),
            R("services", "Сервисы: бот, стрим, лендинг, $ в месяц", "in", "usd", t["services"], NOTE["est"]),
            R("curator", "Кураторы на 1 ученика", "in", "usd", t["curator"], NOTE["curator"]),
            R("tax", "Налог с оборота, %", "in", "pct", t["tax"], NOTE["tax"]),
            R("months", "Длительность сценария, мес", "in", "mon", t["months"]),
            R("costs", "РАСХОДЫ ВСЕГО за 2 запуска", "tot", "usd",
              "={traffic}+{ai_setup}+({ai_sub}+{salary}+{services})*{months}+{m_pct}*{op_plan}+{team}"
              "+{curator}*{students}+{tax}*{net}"),
        ]),
        ("ИТОГ", [
            R("profit", "ПРИБЫЛЬ ЗА 2 ЗАПУСКА", "tot", "usd", "={net}-{costs}"),
            R("ok", "Окупили вложения за 2 месяца?", "tot", "gen", '=IF({profit}>=0,"ДА","НЕТ")'),
            R("payback", "Окупаемость вложений: чистые деньги ÷ расходы", "f", "roi", "=IF({costs}>0,{net}/{costs},0)"),
            R("roas", "Окупаемость трафика: выручка ÷ трафик", "f", "roi", "=IF({traffic}>0,{launch}/{traffic},0)"),
            R("profit_rec", "Прибыль без разовой разработки ИИ", "f", "usd", "={profit}+{ai_setup}"),
            R("start_cash", "Нужно вложить до первых продаж", "f", "usd",
              "={traffic}+{ai_setup}+{team}+{ai_sub}+{salary}+{services}"),
        ]),
    ]


CV1, CV2 = blended_rates()


def build(out_path, addr_path):
    wb = load_workbook(TEMPLATE)
    for name in list(wb.sheetnames):
        if name != SHEET:
            del wb[name]
    ws = wb[SHEET]
    wb.active = 0
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 90), max_col=30):
        for c in row:
            c.value, c.comment = None, None
            c.font, c.fill, c.border, c.alignment, c.number_format = Font(), PatternFill(), Border(), Alignment(), "General"
    for r in list(ws.row_dimensions.keys()):
        ws.row_dimensions[r].height = None

    # ----- раскладка
    top = 7
    s1_plan = Block("S1_plan", stage1_sections(1, plan=True), "B", "C", "D",
                    "ОБЩИЙ ПЛАН — декомпозиция продаж через вебинар", top, title_height=36, sec_fill=C_SEC_B, in_fill=C_IN_B)
    ADDR[("SRC", "total")] = ("I", top + 3)
    ADDR[("SRC", "target")] = ("I", top + 4)
    total = Block("TOT", total_sections(), "P", "Q", "R", "ОБЩИЙ ИТОГ ЗА 2 ЗАПУСКА (2 месяца)", top)
    hdr = max(s1_plan.last_row, total.last_row) + 3
    scen_title = hdr + 4
    s1 = []
    for i, (bid, cols, title, fill) in enumerate([
            ("S1_min", ("B", "C", "D"), "ПЛАН МИНИМУМ", "FFF4CCCC"),
            ("S1_mid", ("G", "H", "I"), "ПЛАН СРЕДНИЙ (= общий план)", "FFFCE5CD"),
            ("S1_max", ("L", "M", "N"), "ПЛАН МАКСИМУМ", "FFB6D7A8")]):
        s1.append(Block(bid, stage1_sections(i, mirror=(i == 1)), *cols, title, scen_title, title_fill=fill,
                        sec_fill=C_SEC_B, in_fill=C_IN_B))
    s2 = []
    for i, (bid, s1bid, cols, title, fill) in enumerate([
            ("S2_min", "S1_min", ("P", "Q", "R"), "ПЛАН МИНИМУМ", "FFF4CCCC"),
            ("S2_mid", "S1_mid", ("U", "V", "W"), "ПЛАН СРЕДНИЙ", "FFFFF2CC"),
            ("S2_max", "S1_max", ("Z", "AA", "AB"), "ПЛАН МАКСИМУМ", "FF00FF00")]):
        s2.append(Block(bid, stage2_sections(i, s1bid), *cols, title, scen_title, title_fill=fill))

    # ----- запись
    for blk in [s1_plan, total] + s1 + s2:
        write_block(ws, blk)

    # суть модели
    ws["B3"] = "Суть модели:"
    ws["B3"].font, ws["B3"].fill, ws["B3"].alignment = mfont(10, True), solid(C_NOTE), CENTER
    merge_box(ws, "C", 3, "AB", 3,
              "Этап 1 — запуск через вебинар: таргет → регистрация в Telegram-бота → вебинар → автооплаты, предоплаты и заказы; "
              "брони и заказы закрывает менеджер, остальных дожимает ИИ-продавец в боте. Часть людей подписывается на ТГ-канал. "
              "Этап 2 — через 2 месяца запуск в ТГ-канале: марафон 2 дня. Итог двух этапов считается по средним сценариям.",
              mfont(10, True), solid(C_NOTE), 48, borders=False)

    # заголовки этапов
    merge_box(ws, "B", 5, "N", 5, "ЭТАП 1. ЗАПУСК ЧЕРЕЗ ВЕБИНАР — холодный трафик с таргета", mfont(13, True), solid(C_TITLE), 30)
    merge_box(ws, "P", 5, "AB", 5, "ИТОГ ДВУХ ЭТАПОВ — сколько денег за 2 запуска и окупили ли (по средним сценариям)",
              mfont(13, True), solid(C_TITLE), 30)
    merge_box(ws, "B", hdr, "N", hdr + 1, "НИЖЕ РАЗНЫЕ СЦЕНАРИИ ПЛАНА ПРОДАЖ — ЭТАП 1 (ВЕБИНАР)", mfont(13, True),
              solid(C_NOTE), borders=False)
    merge_box(ws, "P", hdr, "AB", hdr + 1, "ЭТАП 2. ЗАПУСК ЧЕРЕЗ ТГ-КАНАЛ — марафон 2 дня: разные сценарии плана продаж",
              mfont(13, True), solid(C_TITLE))

    # «От куда?» — декомпозиция регистрации
    merge_box(ws, "H", top, "J", top, "Декомпозиция регистрации", mfont(11, True), solid(C_TITLE), 36)
    for col, text, fill in (("H", "От куда?", "FFFFE599"), ("I", "План", "FFFFE599"), ("J", "Факт", "FFD9EAD3")):
        merge_box(ws, col, top + 1, col, top + 2, text, mfont(11, True), solid(fill))
    first_src = top + 4
    sources = [("Итого", "=SUM(I%d:I%d)" % (first_src, first_src + 10)),
               ("Таргет (Instagram / Facebook)", resolve("=IF({S1_plan:cpl}>0,{S1_plan:traffic}/{S1_plan:cpl},0)", "S1_plan")),
               ("Телеграм канал", 0), ("Бот", 0), ("Емайл", 0), ("Спам тг", 0), ("Спам вотс ап", 0), ("СМС", 0),
               ("ВКонтакте", 0), ("Каналы прошлых мероприятий", 0), ("Отдел продаж", 0), ("Блогеры", 0)]
    for k, (label, val) in enumerate(sources):
        r = top + 3 + k
        for col, v in (("H", label), ("I", val), ("J", None)):
            c = ws[f"{col}{r}"]
            c.value, c.font, c.alignment = v, mfont(11, k <= 1), CENTER
            c.border, c.number_format = brd(l=col == "H", r=True, b=True), FMT["int"]
    ws[f"I{top + 4}"].comment = Comment("Таргет = вложили в трафик ÷ цена регистрации (общий план слева).", "Claude")
    ws[f"H{top + 5}"].comment = Comment("Новая школа: своей базы на старте нет — 0. Если есть канал, бот или блогеры — впиши план.",
                                        "Claude")

    widths = {"A": 2, "B": 14, "C": 46, "D": 16, "E": 3, "F": 3, "G": 14, "H": 46, "I": 16, "J": 14, "K": 3,
              "L": 14, "M": 46, "N": 16, "O": 4, "P": 14, "Q": 50, "R": 16, "S": 3, "T": 3, "U": 14, "V": 50,
              "W": 16, "X": 3, "Y": 3, "Z": 14, "AA": 50, "AB": 16}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    wb.save(out_path)
    Path(addr_path).write_text(json.dumps({f"{b}:{k}": v for (b, k), v in ADDR.items()}, ensure_ascii=False))
    print("ok", out_path, "| hdr", hdr, "| scen title", scen_title, "| plan last", s1_plan.last_row, "| total last", total.last_row)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
