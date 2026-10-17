import copy, json, subprocess, sys
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0, ".")
import single_model as M
from verify_single import verify
RECALC = "/Users/Bel-a/Library/Application Support/Claude/local-agent-mode-sessions/skills-plugin/62ade98b-c20d-435d-bd1e-b573f3f7bff9/c9850139-8eac-49bf-9bdb-4c182fbb50f1/skills/xlsx/scripts/recalc.py"
addr = json.loads(Path("addr.json").read_text())
def setp(p, key, i, val):
    t = list(p[key]); t[i] = val; p[key] = tuple(t)
def run(name, edits, apply):
    wb = load_workbook("single.xlsx"); ws = wb.active
    for key, val in edits.items():
        col, row = addr[key]; ws[f"{col}{row}"] = val
    out = f"stress_{name}.xlsx"; wb.save(out)
    res = json.loads(subprocess.run([sys.executable, RECALC, out, "120"], capture_output=True, text=True).stdout)
    s1, s2, tot = copy.deepcopy(M.S1), copy.deepcopy(M.S2), dict(M.TOT)
    apply(s1, s2, tot)
    probs = verify(out, "addr.json", s1, s2, tot, quiet=True)
    print(f"{name}: errors={res.get('total_errors')} mismatches={probs}")
run("plan_traffic0", {"S1_plan:traffic": 0}, lambda s1, s2, t: setp(s1, "traffic", 1, 0))
run("plan_traffic10k", {"S1_plan:traffic": 10000}, lambda s1, s2, t: setp(s1, "traffic", 1, 10000))
run("max_cpl0", {"S1_max:cpl": 0}, lambda s1, s2, t: setp(s1, "cpl", 2, 0))
run("term1", {"TOT:term": 1}, lambda s1, s2, t: t.update(term=1))
run("months3", {"TOT:months": 3}, lambda s1, s2, t: t.update(months=3))
run("channel3000", {"S2_mid:ch_start": 3000}, lambda s1, s2, t: setp(s2, "ch_start", 1, 3000))
run("no_own", {"TOT:full": 0.70}, lambda s1, s2, t: t.update(full=0.70))
